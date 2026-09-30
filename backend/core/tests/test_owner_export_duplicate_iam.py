"""Focused IAM regressions for owner actions, control exports, and duplication."""

from __future__ import annotations

import io
import uuid

import pytest
from django.contrib.auth.models import Permission
from openpyxl import load_workbook
from rest_framework.exceptions import PermissionDenied
from rest_framework.test import APIClient

from core.models import (
    Actor,
    AppliedControl,
    Evidence,
    EvidenceRevision,
    FilteringLabel,
    Finding,
    FindingsAssessment,
    Policy,
    ReferenceControl,
    RiskMatrix,
    TaskTemplate,
    Terminology,
)
from ebios_rm.models import EbiosRMStudy, Stakeholder
from iam.models import Folder, Role, RoleAssignment, User
from tprm.models import Entity


pytestmark = pytest.mark.django_db


def _domain(name: str) -> Folder:
    return Folder.objects.create(
        name=f"{name}-{uuid.uuid4().hex[:8]}",
        content_type=Folder.ContentType.DOMAIN,
        parent_folder=Folder.get_root_folder(),
    )


def _user(name: str, folder: Folder) -> User:
    user = User.objects.create_user(
        email=f"{name}-{uuid.uuid4().hex}@owner-export.tests"
    )
    user.folder = folder
    user.is_published = False
    user.save(update_fields=["folder", "is_published"])
    return user


def _grant(user: User, folder: Folder, *codenames: str) -> RoleAssignment:
    role = Role.objects.create(
        name=f"Focused IAM role {uuid.uuid4().hex}",
        folder=Folder.get_root_folder(),
    )
    role.permissions.set(Permission.objects.filter(codename__in=codenames))
    assignment = RoleAssignment.objects.create(
        user=user,
        role=role,
        folder=Folder.get_root_folder(),
        is_recursive=False,
    )
    assignment.perimeter_folders.add(folder)
    return assignment


def _client(user: User) -> APIClient:
    client = APIClient()
    client.force_authenticate(user)
    return client


def _actor(name: str, folder: Folder) -> Actor:
    return Actor.objects.get(user=_user(name, folder))


def _control(model, folder: Folder, name: str | None = None):
    return model.objects.create(
        name=name or f"control-{uuid.uuid4().hex}",
        folder=folder,
    )


def _response_rows(response) -> list[dict]:
    payload = response.json()
    return payload.get("results", payload) if isinstance(payload, dict) else payload


@pytest.fixture
def folders():
    Folder._init_root_folder()
    return {
        "visible": _domain("visible"),
        "target": _domain("target"),
        "hidden": _domain("hidden"),
    }


def test_owner_actions_intersect_source_and_actor_iam(folders, monkeypatch):
    monkeypatch.setattr(FindingsAssessment, "upsert_daily_metrics", lambda self: None)
    visible = folders["visible"]
    hidden = folders["hidden"]
    caller = _user("owner-caller", visible)
    _grant(
        caller,
        visible,
        "view_appliedcontrol",
        "view_policy",
        "view_evidence",
        "view_finding",
        "view_user",
    )

    visible_owner = _actor("visible-owner", visible)
    hidden_owner = _actor("hidden-owner", hidden)
    hidden_source_owner = _actor("hidden-source-owner", visible)

    visible_control = _control(AppliedControl, visible)
    hidden_control = _control(AppliedControl, hidden)
    visible_policy = _control(Policy, visible)
    hidden_policy = _control(Policy, hidden)
    visible_evidence = Evidence.objects.create(name="visible-evidence", folder=visible)
    hidden_evidence = Evidence.objects.create(name="hidden-evidence", folder=hidden)
    visible_findings_assessment = FindingsAssessment.objects.create(
        name="visible-findings-assessment", folder=visible
    )
    hidden_findings_assessment = FindingsAssessment.objects.create(
        name="hidden-findings-assessment", folder=hidden
    )
    visible_finding = Finding.objects.create(
        name="visible-finding",
        folder=visible,
        findings_assessment=visible_findings_assessment,
    )
    hidden_finding = Finding.objects.create(
        name="hidden-finding",
        folder=hidden,
        findings_assessment=hidden_findings_assessment,
    )

    for source in (
        visible_control,
        visible_policy,
        visible_evidence,
        visible_finding,
    ):
        source.owner.add(visible_owner, hidden_owner)
    for source in (hidden_control, hidden_policy, hidden_evidence, hidden_finding):
        source.owner.add(hidden_source_owner)

    client = _client(caller)
    for endpoint in ("applied-controls", "policies", "evidences", "findings"):
        response = client.get(f"/api/{endpoint}/owner/")
        assert response.status_code == 200, (endpoint, response.content)
        assert {row["id"] for row in response.json()} == {str(visible_owner.id)}
        rendered = response.content.decode()
        assert hidden_owner.user.email not in rendered
        assert str(hidden_owner.id) not in rendered
        assert hidden_source_owner.user.email not in rendered
        assert str(hidden_source_owner.id) not in rendered


def test_budget_analytics_omits_hidden_owner_and_folder_buckets(folders):
    visible = folders["visible"]
    hidden = folders["hidden"]
    caller = _user("budget-caller", visible)
    _grant(
        caller,
        visible,
        "view_appliedcontrol",
        "view_folder",
        "view_user",
    )
    # The caller may aggregate controls in this second domain, but has no
    # authority to identify its Folder or Actor dimensions.
    _grant(caller, hidden, "view_appliedcontrol")

    visible_owner = _actor("budget-visible-owner", visible)
    hidden_owner = _actor("budget-hidden-owner", hidden)
    visible_control = _control(AppliedControl, visible)
    hidden_control = _control(AppliedControl, hidden)
    visible_control.owner.add(visible_owner)
    hidden_control.owner.add(hidden_owner)

    response = _client(caller).get("/api/applied-controls/analytics/")

    assert response.status_code == 200, response.content
    payload = response.json()
    assert payload["count"] == 2
    assert {bucket["key"] for bucket in payload["top_owners"]} == {
        str(visible_owner.id)
    }
    assert {bucket["key"] for bucket in payload["top_folders"]} == {str(visible.id)}
    rendered = response.content.decode()
    assert str(hidden_owner.id) not in rendered
    assert hidden_owner.user.email not in rendered
    assert str(hidden.id) not in rendered
    assert hidden.name not in rendered


def _export_text(response, action: str) -> str:
    if action == "export_csv":
        return response.content.decode("utf-8-sig")
    workbook = load_workbook(io.BytesIO(response.content), read_only=True)
    return "\n".join(
        str(cell)
        for worksheet in workbook.worksheets
        for row in worksheet.iter_rows(values_only=True)
        for cell in row
        if cell is not None
    )


@pytest.mark.parametrize(
    "model,endpoint,view_codename",
    (
        (AppliedControl, "applied-controls", "view_appliedcontrol"),
        (Policy, "policies", "view_policy"),
    ),
)
@pytest.mark.parametrize("action", ("export_csv", "export_xlsx"))
def test_control_exports_project_related_metadata_through_request_iam(
    folders, model, endpoint, view_codename, action
):
    visible = folders["visible"]
    hidden = folders["hidden"]
    caller = _user("export-caller", visible)
    _grant(
        caller,
        visible,
        view_codename,
        "view_user",
        "view_filteringlabel",
        "view_evidence",
        "view_evidencerevision",
    )

    control = _control(model, visible, f"export-control-{uuid.uuid4().hex}")
    visible_owner = _actor("export-visible-owner", visible)
    hidden_owner = _actor("export-hidden-owner", hidden)
    control.owner.add(visible_owner, hidden_owner)

    visible_label = FilteringLabel.objects.create(
        label=f"visible-{uuid.uuid4().hex[:8]}", folder=visible
    )
    hidden_label = FilteringLabel.objects.create(
        label=f"hidden-{uuid.uuid4().hex[:8]}", folder=hidden
    )
    control.filtering_labels.add(visible_label, hidden_label)

    visible_evidence = Evidence.objects.create(
        name=f"visible-evidence-{uuid.uuid4().hex}", folder=visible
    )
    hidden_evidence = Evidence.objects.create(
        name=f"hidden-evidence-{uuid.uuid4().hex}", folder=hidden
    )
    visible_filename = f"visible-{uuid.uuid4().hex}.txt"
    hidden_revision_filename = f"hidden-revision-{uuid.uuid4().hex}.txt"
    hidden_evidence_filename = f"hidden-evidence-{uuid.uuid4().hex}.txt"
    EvidenceRevision.objects.create(
        evidence=visible_evidence,
        folder=visible,
        version=1,
        attachment=f"evidence/{visible_filename}",
    )
    hidden_revision = EvidenceRevision.objects.create(
        evidence=visible_evidence,
        folder=visible,
        version=2,
        attachment=f"evidence/{hidden_revision_filename}",
    )
    EvidenceRevision.objects.filter(id=hidden_revision.id).update(folder=hidden)
    EvidenceRevision.objects.create(
        evidence=hidden_evidence,
        folder=hidden,
        version=1,
        attachment=f"evidence/{hidden_evidence_filename}",
    )
    control.evidences.add(visible_evidence, hidden_evidence)

    response = _client(caller).get(f"/api/{endpoint}/{action}/")

    assert response.status_code == 200, response.content
    rendered = _export_text(response, action)
    assert visible_owner.user.email in rendered
    assert visible_label.label in rendered
    assert visible_evidence.name in rendered
    assert visible_filename in rendered
    assert hidden_owner.user.email not in rendered
    assert str(hidden_owner.id) not in rendered
    assert hidden_label.label not in rendered
    assert str(hidden_label.id) not in rendered
    assert hidden_evidence.name not in rendered
    assert str(hidden_evidence.id) not in rendered
    assert hidden_revision_filename not in rendered
    assert hidden_evidence_filename not in rendered


@pytest.mark.parametrize(
    "model,endpoint,view_codename",
    (
        (AppliedControl, "applied-controls", "view_appliedcontrol"),
        (Policy, "policies", "view_policy"),
    ),
)
@pytest.mark.parametrize("action", ("export_csv", "export_xlsx"))
def test_control_exports_mask_hidden_folder_and_reference_control(
    folders, model, endpoint, view_codename, action
):
    visible = folders["visible"]
    hidden = folders["hidden"]
    caller = _user("export-singular-caller", visible)
    _grant(caller, visible, view_codename, "view_folder", "view_referencecontrol")
    # The control row is deliberately readable while its containing folder is
    # not.  Singular related metadata needs its own IAM projection.
    _grant(caller, hidden, view_codename)

    visible_reference = ReferenceControl.objects.create(
        name=f"visible-reference-{uuid.uuid4().hex}",
        ref_id=f"VISIBLE-{uuid.uuid4().hex[:8]}",
        urn=f"urn:test:visible-reference:{uuid.uuid4().hex}",
        folder=visible,
    )
    hidden_reference = ReferenceControl.objects.create(
        name=f"hidden-reference-{uuid.uuid4().hex}",
        ref_id=f"HIDDEN-{uuid.uuid4().hex[:8]}",
        urn=f"urn:test:hidden-reference:{uuid.uuid4().hex}",
        folder=hidden,
    )
    visible_control = _control(
        model, visible, f"visible-export-control-{uuid.uuid4().hex}"
    )
    visible_control.reference_control = visible_reference
    visible_control.save(update_fields=["reference_control"])
    hidden_control = _control(
        model, hidden, f"hidden-folder-export-control-{uuid.uuid4().hex}"
    )
    hidden_control.reference_control = hidden_reference
    hidden_control.save(update_fields=["reference_control"])

    response = _client(caller).get(f"/api/{endpoint}/{action}/")

    assert response.status_code == 200, response.content
    rendered = _export_text(response, action)
    assert visible_control.name in rendered
    assert visible.name in rendered
    assert visible_reference.name in rendered
    assert visible_reference.ref_id in rendered
    assert hidden_control.name in rendered
    assert hidden.name not in rendered
    assert hidden_reference.name not in rendered
    assert hidden_reference.ref_id not in rendered


@pytest.mark.parametrize(
    "model,endpoint,view_codename",
    (
        (AppliedControl, "applied-controls", "view_appliedcontrol"),
        (Policy, "policies", "view_policy"),
    ),
)
def test_control_assignment_projection_and_filter_ignore_hidden_actor_owners(
    folders, model, endpoint, view_codename
):
    visible = folders["visible"]
    hidden = folders["hidden"]
    caller = _user("assignment-projection-caller", visible)
    _grant(caller, visible, view_codename, "view_user")

    visible_owner = _actor("assignment-visible-owner", visible)
    hidden_owner = _actor("assignment-hidden-owner", hidden)
    visible_owned = _control(model, visible, f"visible-owned-{uuid.uuid4().hex}")
    hidden_only_owned = _control(
        model, visible, f"hidden-only-owned-{uuid.uuid4().hex}"
    )
    visible_owned.owner.add(visible_owner)
    hidden_only_owned.owner.add(hidden_owner)
    client = _client(caller)

    list_response = client.get(f"/api/{endpoint}/")
    detail_response = client.get(f"/api/{endpoint}/{hidden_only_owned.id}/")
    full_response = client.get(f"/api/{endpoint}/full/")

    assert list_response.status_code == 200, list_response.content
    assert detail_response.status_code == 200, detail_response.content
    assert full_response.status_code == 200, full_response.content

    list_rows = {row["id"]: row for row in _response_rows(list_response)}
    full_rows = {row["id"]: row for row in _response_rows(full_response)}
    for rows in (list_rows, full_rows):
        assert rows[str(visible_owned.id)]["is_assigned"] is True
        assert rows[str(hidden_only_owned.id)]["is_assigned"] is False
        assert str(hidden_owner.id) not in str(rows[str(hidden_only_owned.id)]["owner"])

    detail_payload = detail_response.json()
    assert detail_payload["is_assigned"] is False
    assert str(hidden_owner.id) not in str(detail_payload["owner"])
    assert hidden_owner.user.email not in detail_response.content.decode()

    assigned_response = client.get(f"/api/{endpoint}/", {"is_assigned": "true"})
    unassigned_response = client.get(f"/api/{endpoint}/", {"is_assigned": "false"})

    assert assigned_response.status_code == 200, assigned_response.content
    assert unassigned_response.status_code == 200, unassigned_response.content
    assigned_ids = {row["id"] for row in _response_rows(assigned_response)}
    unassigned_ids = {row["id"] for row in _response_rows(unassigned_response)}
    assert str(visible_owned.id) in assigned_ids
    assert str(hidden_only_owned.id) not in assigned_ids
    assert str(hidden_only_owned.id) in unassigned_ids
    assert str(visible_owned.id) not in unassigned_ids


@pytest.mark.parametrize(
    "model,endpoint,view_codename",
    (
        (AppliedControl, "applied-controls", "view_appliedcontrol"),
        (Policy, "policies", "view_policy"),
    ),
)
@pytest.mark.parametrize("action", ("export_csv", "export_xlsx"))
def test_control_export_is_assigned_filter_uses_only_visible_owners(
    folders, model, endpoint, view_codename, action
):
    visible = folders["visible"]
    hidden = folders["hidden"]
    caller = _user("assignment-export-caller", visible)
    _grant(caller, visible, view_codename, "view_user")

    visible_owner = _actor("assignment-export-visible-owner", visible)
    hidden_owner = _actor("assignment-export-hidden-owner", hidden)
    visible_owned = _control(model, visible, f"export-visible-owned-{uuid.uuid4().hex}")
    hidden_only_owned = _control(
        model, visible, f"export-hidden-only-owned-{uuid.uuid4().hex}"
    )
    visible_owned.owner.add(visible_owner)
    hidden_only_owned.owner.add(hidden_owner)
    client = _client(caller)

    assigned_response = client.get(
        f"/api/{endpoint}/{action}/", {"is_assigned": "true"}
    )
    unassigned_response = client.get(
        f"/api/{endpoint}/{action}/", {"is_assigned": "false"}
    )

    assert assigned_response.status_code == 200, assigned_response.content
    assert unassigned_response.status_code == 200, unassigned_response.content
    assigned_rendered = _export_text(assigned_response, action)
    unassigned_rendered = _export_text(unassigned_response, action)
    assert visible_owned.name in assigned_rendered
    assert hidden_only_owned.name not in assigned_rendered
    assert visible_owner.user.email in assigned_rendered
    assert hidden_owner.user.email not in assigned_rendered
    assert hidden_only_owned.name in unassigned_rendered
    assert visible_owned.name not in unassigned_rendered
    assert hidden_owner.user.email not in unassigned_rendered
    assert str(hidden_owner.id) not in unassigned_rendered


@pytest.mark.parametrize(
    "model,endpoint,view_codename",
    (
        (AppliedControl, "applied-controls", "view_appliedcontrol"),
        (Policy, "policies", "view_policy"),
    ),
)
def test_related_filter_hidden_and_missing_operands_share_generic_response(
    folders, model, endpoint, view_codename
):
    visible = folders["visible"]
    hidden = folders["hidden"]
    caller = _user("filter-oracle-caller", visible)
    _grant(
        caller,
        visible,
        view_codename,
        "view_folder",
        "view_referencecontrol",
        "view_evidence",
        "view_user",
    )
    # The caller can read a control stored in this folder, but cannot read the
    # Folder row itself. The filter operand must use the relation's own IAM.
    _grant(caller, hidden, view_codename)

    hidden_folder_control = _control(
        model, hidden, f"hidden-folder-linked-{uuid.uuid4().hex}"
    )
    linked_control = _control(model, visible, f"hidden-relations-{uuid.uuid4().hex}")
    hidden_reference = ReferenceControl.objects.create(
        name=f"hidden-filter-reference-{uuid.uuid4().hex}",
        ref_id=f"HIDDEN-FILTER-{uuid.uuid4().hex[:8]}",
        urn=f"urn:test:hidden-filter-reference:{uuid.uuid4().hex}",
        folder=hidden,
    )
    linked_control.reference_control = hidden_reference
    linked_control.save(update_fields=["reference_control"])
    hidden_evidence = Evidence.objects.create(
        name=f"hidden-filter-evidence-{uuid.uuid4().hex}", folder=hidden
    )
    hidden_owner = _actor("hidden-filter-owner", hidden)
    linked_control.evidences.add(hidden_evidence)
    linked_control.owner.add(hidden_owner)
    client = _client(caller)

    operands = {
        "folder": hidden.id,
        "reference_control": hidden_reference.id,
        "evidences": hidden_evidence.id,
        "owner": hidden_owner.id,
    }
    association_tokens = (
        str(hidden_folder_control.id),
        hidden_folder_control.name,
        str(linked_control.id),
        linked_control.name,
        hidden_reference.name,
        hidden_reference.ref_id,
        hidden_evidence.name,
        hidden_owner.user.email,
    )

    for filter_name, hidden_id in operands.items():
        missing_id = uuid.uuid4()
        hidden_response = client.get(f"/api/{endpoint}/", {filter_name: str(hidden_id)})
        missing_response = client.get(
            f"/api/{endpoint}/", {filter_name: str(missing_id)}
        )

        assert hidden_response.status_code == missing_response.status_code == 400
        # Django's generic ModelChoice error may repeat the caller-supplied
        # operand. Normalize that value before comparing the response contract.
        hidden_body = hidden_response.content.decode().replace(
            str(hidden_id), "<operand>"
        )
        missing_body = missing_response.content.decode().replace(
            str(missing_id), "<operand>"
        )
        assert hidden_body == missing_body
        for token in association_tokens:
            assert token not in hidden_body


def _hidden_stakeholder(folder: Folder) -> Stakeholder:
    matrix = RiskMatrix.objects.create(
        name=f"hidden-stakeholder-matrix-{uuid.uuid4().hex}",
        urn=f"urn:test:hidden-stakeholder-matrix:{uuid.uuid4().hex}",
        folder=folder,
        json_definition={
            "probability": [{"name": "possible"}],
            "impact": [{"name": "limited"}],
            "risk": [{"name": "low"}],
            "grid": [[0]],
        },
    )
    entity = Entity.objects.create(
        name=f"hidden-stakeholder-entity-{uuid.uuid4().hex}", folder=folder
    )
    study = EbiosRMStudy.objects.create(
        name=f"hidden-stakeholder-study-{uuid.uuid4().hex}",
        folder=folder,
        risk_matrix=matrix,
        reference_entity=entity,
    )
    category = Terminology.objects.create(
        name=f"hidden-stakeholder-category-{uuid.uuid4().hex}",
        folder=folder,
        field_path=Terminology.FieldPath.ENTITY_RELATIONSHIP,
    )
    return Stakeholder.objects.create(
        folder=folder,
        ebios_rm_study=study,
        entity=entity,
        category=category,
    )


@pytest.mark.parametrize(
    "model,endpoint,view_codename,change_codename",
    (
        (
            AppliedControl,
            "applied-controls",
            "view_appliedcontrol",
            "change_appliedcontrol",
        ),
        (Policy, "policies", "view_policy", "change_policy"),
    ),
)
def test_unrelated_patch_projects_hidden_relations_from_response(
    folders,
    monkeypatch,
    model,
    endpoint,
    view_codename,
    change_codename,
):
    monkeypatch.setattr(FindingsAssessment, "upsert_daily_metrics", lambda self: None)
    visible = folders["visible"]
    hidden = folders["hidden"]
    caller = _user("mutation-projection-caller", visible)
    _grant(
        caller,
        visible,
        view_codename,
        change_codename,
        "view_finding",
        "view_stakeholder",
        "view_tasktemplate",
        "view_user",
        "view_evidence",
    )
    control = _control(model, visible, f"mutation-projection-{uuid.uuid4().hex}")

    visible_assessment = FindingsAssessment.objects.create(
        name=f"visible-mutation-assessment-{uuid.uuid4().hex}", folder=visible
    )
    hidden_assessment = FindingsAssessment.objects.create(
        name=f"hidden-mutation-assessment-{uuid.uuid4().hex}", folder=hidden
    )
    visible_finding = Finding.objects.create(
        name=f"visible-mutation-finding-{uuid.uuid4().hex}",
        folder=visible,
        findings_assessment=visible_assessment,
    )
    hidden_finding = Finding.objects.create(
        name=f"hidden-mutation-finding-{uuid.uuid4().hex}",
        folder=hidden,
        findings_assessment=hidden_assessment,
    )
    hidden_stakeholder = _hidden_stakeholder(hidden)
    hidden_task = TaskTemplate.objects.create(
        name=f"hidden-mutation-task-{uuid.uuid4().hex}", folder=hidden
    )
    hidden_owner = _actor("hidden-mutation-owner", hidden)
    hidden_evidence = Evidence.objects.create(
        name=f"hidden-mutation-evidence-{uuid.uuid4().hex}", folder=hidden
    )
    control.findings.add(visible_finding, hidden_finding)
    hidden_stakeholder.applied_controls.add(control)
    hidden_task.applied_controls.add(control)
    control.owner.add(hidden_owner)
    control.evidences.add(hidden_evidence)

    updated_description = f"unrelated-update-{uuid.uuid4().hex}"
    response = _client(caller).patch(
        f"/api/{endpoint}/{control.id}/",
        {"description": updated_description},
        format="json",
    )

    assert response.status_code == 200, response.content
    payload = response.json()
    assert payload["description"] == updated_description
    assert isinstance(payload["findings"], list)
    assert payload["findings"] == [str(visible_finding.id)]
    rendered = response.content.decode()
    for hidden_object in (
        hidden_finding,
        hidden_stakeholder,
        hidden_task,
        hidden_owner,
        hidden_evidence,
    ):
        assert str(hidden_object.id) not in rendered
    assert hidden_finding.name not in rendered
    assert hidden_task.name not in rendered
    assert hidden_owner.user.email not in rendered
    assert hidden_evidence.name not in rendered


@pytest.mark.parametrize(
    "model,endpoint,view_codename,add_codename",
    (
        (
            AppliedControl,
            "applied-controls",
            "view_appliedcontrol",
            "add_appliedcontrol",
        ),
        (Policy, "policies", "view_policy", "add_policy"),
    ),
)
def test_duplicate_copies_only_visible_owner_and_label(
    folders, model, endpoint, view_codename, add_codename
):
    visible = folders["visible"]
    target = folders["target"]
    hidden = folders["hidden"]
    caller = _user("duplicate-caller", visible)
    _grant(
        caller,
        visible,
        view_codename,
        "view_user",
        "view_filteringlabel",
    )
    _grant(caller, target, add_codename)
    source = _control(model, visible)
    visible_owner = _actor("duplicate-visible-owner", visible)
    hidden_owner = _actor("duplicate-hidden-owner", hidden)
    visible_label = FilteringLabel.objects.create(
        label=f"visible-{uuid.uuid4().hex[:8]}", folder=visible
    )
    hidden_label = FilteringLabel.objects.create(
        label=f"hidden-{uuid.uuid4().hex[:8]}", folder=hidden
    )
    source.owner.add(visible_owner, hidden_owner)
    source.filtering_labels.add(visible_label, hidden_label)
    duplicate_name = f"duplicate-{uuid.uuid4().hex}"

    response = _client(caller).post(
        f"/api/{endpoint}/{source.id}/duplicate/",
        {
            "name": duplicate_name,
            "description": "",
            "folder": str(target.id),
            "duplicate_evidences": False,
        },
        format="json",
    )

    assert response.status_code == 200, response.content
    duplicate = model.objects.get(name=duplicate_name)
    assert set(duplicate.owner.values_list("id", flat=True)) == {visible_owner.id}
    assert set(duplicate.filtering_labels.values_list("id", flat=True)) == {
        visible_label.id
    }
    rendered = response.content.decode()
    assert str(visible_owner.id) in rendered
    assert str(visible_label.id) in rendered
    assert str(hidden_owner.id) not in rendered
    assert hidden_owner.user.email not in rendered
    assert str(hidden_label.id) not in rendered
    assert hidden_label.label not in rendered


@pytest.mark.parametrize(
    "model,endpoint,view_codename,add_codename",
    (
        (
            AppliedControl,
            "applied-controls",
            "view_appliedcontrol",
            "add_appliedcontrol",
        ),
        (Policy, "policies", "view_policy", "add_policy"),
    ),
)
def test_duplicate_fails_closed_for_hidden_reference_control(
    folders, model, endpoint, view_codename, add_codename
):
    visible = folders["visible"]
    target = folders["target"]
    hidden = folders["hidden"]
    caller = _user("duplicate-hidden-reference", visible)
    _grant(caller, visible, view_codename)
    _grant(caller, target, add_codename)
    hidden_reference = ReferenceControl.objects.create(
        name=f"hidden-duplicate-reference-{uuid.uuid4().hex}",
        urn=f"urn:test:hidden-duplicate-reference:{uuid.uuid4().hex}",
        folder=hidden,
    )
    source = _control(model, visible)
    source.reference_control = hidden_reference
    source.save(update_fields=["reference_control"])
    duplicate_name = f"hidden-reference-duplicate-{uuid.uuid4().hex}"

    response = _client(caller).post(
        f"/api/{endpoint}/{source.id}/duplicate/",
        {
            "name": duplicate_name,
            "description": "",
            "folder": str(target.id),
            "duplicate_evidences": False,
        },
        format="json",
    )

    assert response.status_code == 403, response.content
    assert not model.objects.filter(name=duplicate_name).exists()
    rendered = response.content.decode()
    assert str(hidden_reference.id) not in rendered
    assert hidden_reference.name not in rendered


@pytest.mark.parametrize(
    "model,endpoint,view_codename",
    (
        (AppliedControl, "applied-controls", "view_appliedcontrol"),
        (Policy, "policies", "view_policy"),
    ),
)
def test_duplicate_hidden_and_missing_target_folder_share_denial(
    folders, model, endpoint, view_codename
):
    visible = folders["visible"]
    hidden = folders["hidden"]
    caller = _user("duplicate-target-oracle", visible)
    _grant(caller, visible, view_codename)
    source = _control(model, visible)
    client = _client(caller)

    responses = []
    for index, target_id in enumerate((hidden.id, uuid.uuid4())):
        duplicate_name = f"target-oracle-{index}-{uuid.uuid4().hex}"
        response = client.post(
            f"/api/{endpoint}/{source.id}/duplicate/",
            {
                "name": duplicate_name,
                "description": "",
                "folder": str(target_id),
                "duplicate_evidences": False,
            },
            format="json",
        )
        assert not model.objects.filter(name=duplicate_name).exists()
        responses.append(response)

    assert [response.status_code for response in responses] == [403, 403]
    assert responses[0].json() == responses[1].json()
    assert str(hidden.id) not in responses[0].content.decode()


@pytest.mark.parametrize(
    "model,endpoint,view_codename,add_codename",
    (
        (
            AppliedControl,
            "applied-controls",
            "view_appliedcontrol",
            "add_appliedcontrol",
        ),
        (Policy, "policies", "view_policy", "add_policy"),
    ),
)
@pytest.mark.parametrize("failure", ("hidden_source", "missing_add_permission"))
def test_duplicate_evidences_fails_closed_before_creating_control(
    folders, model, endpoint, view_codename, add_codename, failure
):
    visible = folders["visible"]
    target = folders["target"]
    hidden = folders["hidden"]
    caller = _user("duplicate-denied", visible)
    _grant(caller, visible, view_codename, "view_evidence")
    _grant(caller, target, add_codename, "view_evidence")
    source = _control(model, visible)
    evidence_folder = hidden if failure == "hidden_source" else visible
    evidence = Evidence.objects.create(
        name=f"source-evidence-{uuid.uuid4().hex}", folder=evidence_folder
    )
    source.evidences.add(evidence)
    duplicate_name = f"denied-duplicate-{uuid.uuid4().hex}"

    response = _client(caller).post(
        f"/api/{endpoint}/{source.id}/duplicate/",
        {
            "name": duplicate_name,
            "description": "",
            "folder": str(target.id),
            "duplicate_evidences": True,
        },
        format="json",
    )

    assert response.status_code == 403, response.content
    assert not model.objects.filter(name=duplicate_name).exists()
    assert not Evidence.objects.filter(name=evidence.name, folder=target).exists()


def test_duplicate_does_not_link_hidden_same_name_target(folders):
    visible = folders["visible"]
    target = folders["target"]
    caller = _user("duplicate-hidden-target", visible)
    _grant(caller, visible, "view_appliedcontrol", "view_evidence")
    _grant(caller, target, "add_appliedcontrol", "add_evidence")
    source = _control(AppliedControl, visible)
    evidence_name = f"same-name-{uuid.uuid4().hex}"
    source_evidence = Evidence.objects.create(name=evidence_name, folder=visible)
    hidden_target = Evidence.objects.create(name=evidence_name.upper(), folder=target)
    source.evidences.add(source_evidence)
    duplicate_name = f"safe-duplicate-{uuid.uuid4().hex}"

    response = _client(caller).post(
        f"/api/applied-controls/{source.id}/duplicate/",
        {
            "name": duplicate_name,
            "description": "",
            "folder": str(target.id),
            "duplicate_evidences": True,
        },
        format="json",
    )

    assert response.status_code == 403, response.content
    assert not AppliedControl.objects.filter(name=duplicate_name).exists()
    assert (
        Evidence.objects.filter(name__iexact=evidence_name, folder=target).count() == 1
    )
    assert str(hidden_target.id) not in response.content.decode()


def test_duplicate_can_link_visible_same_name_target_without_add_evidence(folders):
    visible = folders["visible"]
    target = folders["target"]
    caller = _user("duplicate-visible-target", visible)
    _grant(caller, visible, "view_appliedcontrol", "view_evidence")
    _grant(
        caller,
        target,
        "add_appliedcontrol",
        "view_appliedcontrol",
        "view_evidence",
    )
    source = _control(AppliedControl, visible)
    evidence_name = f"same-name-{uuid.uuid4().hex}"
    source_evidence = Evidence.objects.create(name=evidence_name, folder=visible)
    visible_target = Evidence.objects.create(name=evidence_name.upper(), folder=target)
    source.evidences.add(source_evidence)

    response = _client(caller).post(
        f"/api/applied-controls/{source.id}/duplicate/",
        {
            "name": f"linked-duplicate-{uuid.uuid4().hex}",
            "description": "",
            "folder": str(target.id),
            "duplicate_evidences": True,
        },
        format="json",
    )

    assert response.status_code == 200, response.content
    duplicate = AppliedControl.objects.get(id=response.json()["results"]["id"])
    assert set(duplicate.evidences.values_list("id", flat=True)) == {visible_target.id}
    assert (
        Evidence.objects.filter(name__iexact=evidence_name, folder=target).count() == 1
    )
    assert response.json()["results"]["evidences"][0]["id"] == str(visible_target.id)


def test_duplicate_evidence_failure_rolls_back_control_and_prior_clone(
    folders, monkeypatch
):
    visible = folders["visible"]
    target = folders["target"]
    caller = _user("duplicate-rollback", visible)
    _grant(caller, visible, "view_appliedcontrol", "view_evidence")
    _grant(caller, target, "add_appliedcontrol", "add_evidence")
    source = _control(AppliedControl, visible)
    evidence_names = [
        f"rollback-a-{uuid.uuid4().hex}",
        f"rollback-b-{uuid.uuid4().hex}",
    ]
    source.evidences.add(
        *[Evidence.objects.create(name=name, folder=visible) for name in evidence_names]
    )
    duplicate_name = f"rollback-control-{uuid.uuid4().hex}"
    manager_class = Evidence.objects.__class__
    original_create = manager_class.create
    evidence_create_count = 0

    def fail_second_evidence_create(manager, *args, **kwargs):
        nonlocal evidence_create_count
        if manager.model is Evidence:
            evidence_create_count += 1
            if evidence_create_count == 2:
                raise PermissionDenied("forced evidence clone failure")
        return original_create(manager, *args, **kwargs)

    monkeypatch.setattr(manager_class, "create", fail_second_evidence_create)

    response = _client(caller).post(
        f"/api/applied-controls/{source.id}/duplicate/",
        {
            "name": duplicate_name,
            "description": "",
            "folder": str(target.id),
            "duplicate_evidences": True,
        },
        format="json",
    )

    assert response.status_code == 403, response.content
    assert evidence_create_count == 2
    assert not AppliedControl.objects.filter(name=duplicate_name).exists()
    assert not Evidence.objects.filter(name__in=evidence_names, folder=target).exists()
