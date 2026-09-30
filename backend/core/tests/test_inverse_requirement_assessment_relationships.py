"""Focused authority tests for inverse RequirementAssessment relationship writes."""

from __future__ import annotations

import io
import uuid

import pytest
from django.contrib.auth.models import Permission
from django.contrib.contenttypes.models import ContentType
from django.db import connection
from django.db.models import QuerySet
from django.test.utils import CaptureQueriesContext
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework import serializers
from rest_framework.exceptions import (
    NotAuthenticated,
    PermissionDenied,
    ValidationError,
)
from rest_framework.test import APIClient, APIRequestFactory

from core.models import (
    Actor,
    AppliedControl,
    Assessment,
    ComplianceAssessment,
    Evidence,
    EvidenceRevision,
    Finding,
    FindingsAssessment,
    Framework,
    Perimeter,
    Policy,
    ReferenceControl,
    RequirementAssessment,
    RequirementAssignment,
    RequirementNode,
    RiskAssessment,
    RiskMatrix,
    RiskScenario,
    SecurityException,
    TaskTemplate,
    Team,
)
from core.requirement_assessment_relationships import (
    BATCH_RELATION_OPERATION_CONTEXT_KEY,
    RequirementAssessmentRelationshipAuthorityMixin,
    RequirementAssessmentRelationshipConflict,
    RequirementAssessmentRelationshipProjectionListSerializer,
    RequirementAssessmentRelationshipProjectionMixin,
    _assignment_authority_snapshot,
)
from core.serializer_fields import FieldsRelatedField
from core.serializers import (
    AppliedControlBulkReadSerializer,
    AppliedControlListSerializer,
    AppliedControlReadSerializer,
    AppliedControlRequestProjectionListSerializer,
    AppliedControlWriteSerializer,
    EvidenceReadSerializer,
    EvidenceWriteSerializer,
    PolicyReadSerializer,
    PolicyWriteSerializer,
    RiskAssessmentActionPlanSerializer,
    SecurityExceptionReadSerializer,
    SecurityExceptionWriteSerializer,
    SerializerFactory,
)
from custom_fields.models import CustomFieldDefinition, CustomFieldValue, FieldType
from iam.models import Folder, Role, RoleAssignment, User
from integrations.models import IntegrationConfiguration, IntegrationProvider

pytestmark = pytest.mark.django_db


TARGETS = (
    ("evidences", Evidence, EvidenceWriteSerializer),
    ("applied_controls", AppliedControl, AppliedControlWriteSerializer),
    (
        "security_exceptions",
        SecurityException,
        SecurityExceptionWriteSerializer,
    ),
    ("applied_controls", Policy, PolicyWriteSerializer),
)
CONCRETE_TARGETS = TARGETS[:3]
READ_SERIALIZERS = {
    Evidence: EvidenceReadSerializer,
    AppliedControl: AppliedControlReadSerializer,
    SecurityException: SecurityExceptionReadSerializer,
    Policy: PolicyReadSerializer,
}
ENDPOINTS = {
    Evidence: "evidences",
    AppliedControl: "applied-controls",
    SecurityException: "security-exceptions",
    Policy: "policies",
}


class _ProjectionOnlyEvidenceSerializer(
    RequirementAssessmentRelationshipProjectionMixin,
    serializers.ModelSerializer,
):
    requirement_assessments = FieldsRelatedField(many=True)

    class Meta:
        model = Evidence
        fields = ["id", "requirement_assessments"]
        list_serializer_class = (
            RequirementAssessmentRelationshipProjectionListSerializer
        )


def _grant_permissions(user: User, folder: Folder, *codenames: str) -> None:
    role = Role.objects.create(
        name=f"Inverse RA permission scope {uuid.uuid4().hex}",
        folder=Folder.get_root_folder(),
    )
    role.permissions.set(Permission.objects.filter(codename__in=codenames))
    role_assignment = RoleAssignment.objects.create(
        user=user,
        role=role,
        folder=Folder.get_root_folder(),
        is_recursive=True,
    )
    role_assignment.perimeter_folders.add(folder)


def _grant(user: User, folder: Folder) -> None:
    codenames = (
        "view_complianceassessment",
        "view_framework",
        "view_requirementnode",
        "view_requirementassessment",
        "change_requirementassessment",
        "view_requirementassignment",
        "view_evidence",
        "add_evidence",
        "change_evidence",
        "view_appliedcontrol",
        "add_appliedcontrol",
        "change_appliedcontrol",
        "view_securityexception",
        "add_securityexception",
        "change_securityexception",
        "view_policy",
        "add_policy",
        "change_policy",
        "view_finding",
        "view_riskassessment",
        "view_riskscenario",
        "view_tasktemplate",
        "view_user",
    )
    _grant_permissions(user, folder, *codenames)


def _request_for(user: User | None):
    request = APIRequestFactory().patch("/", {}, format="json")
    request.user = user
    return request


def _target(model, folder: Folder):
    return model.objects.create(
        name=f"inverse-ra-target-{uuid.uuid4().hex}",
        folder=folder,
    )


def _serializer(serializer_class, *, target, user, requested, operation=None, **data):
    context = {"request": _request_for(user)}
    if operation is not None:
        context[BATCH_RELATION_OPERATION_CONTEXT_KEY] = operation
    return serializer_class(
        target,
        data={
            "requirement_assessments": [str(row.id) for row in requested],
            **data,
        },
        partial=True,
        context=context,
    )


def _represented_ids(values):
    return {
        str(value.get("id") if isinstance(value, dict) else value) for value in values
    }


@pytest.fixture
def relationship_world():
    Folder._init_root_folder()
    root = Folder.get_root_folder()
    folder = Folder.objects.create(
        name=f"inverse-ra-domain-{uuid.uuid4().hex[:8]}",
        content_type=Folder.ContentType.DOMAIN,
        parent_folder=root,
    )
    framework = Framework.objects.create(
        name="Inverse relationship framework",
        urn=f"urn:test:inverse-ra:{uuid.uuid4().hex}",
        ref_id="INVERSE-RA",
        folder=folder,
        min_score=0,
        max_score=4,
    )
    requirements = [
        RequirementNode.objects.create(
            name=name,
            urn=f"{framework.urn}:{suffix}",
            ref_id=suffix.upper(),
            framework=framework,
            folder=folder,
            assessable=True,
        )
        for name, suffix in (
            ("Assigned A", "assigned-a"),
            ("Assigned B", "assigned-b"),
            ("Unassigned", "unassigned"),
        )
    ]
    assessment = ComplianceAssessment.objects.create(
        name="Inverse relationship audit",
        ref_id="INVERSE-RA-AUDIT",
        framework=framework,
        folder=folder,
        perimeter=Perimeter.objects.create(name="Inverse perimeter", folder=folder),
        min_score=0,
        max_score=4,
        status=Assessment.Status.IN_PROGRESS,
        field_visibility={
            field_name: {"auditor": "edit", "respondent": "edit"}
            for field_name, _model, _serializer_class in TARGETS
        },
    )
    assigned_a, assigned_b, unassigned = [
        RequirementAssessment.objects.create(
            compliance_assessment=assessment,
            requirement=requirement,
            folder=folder,
        )
        for requirement in requirements
    ]

    respondent = User.objects.create_user(
        email=f"inverse-ra-{uuid.uuid4().hex}@tests.invalid"
    )
    _grant(respondent, folder)
    actor, _created = Actor.objects.get_or_create(user=respondent)
    assignment = RequirementAssignment.objects.create(
        compliance_assessment=assessment,
        folder=folder,
        status=RequirementAssignment.Status.IN_PROGRESS,
    )
    assignment.actor.add(actor)
    assignment.requirement_assessments.add(assigned_a, assigned_b)

    return {
        "folder": folder,
        "assessment": assessment,
        "assignment": assignment,
        "assigned_a": assigned_a,
        "assigned_b": assigned_b,
        "unassigned": unassigned,
        "respondent": respondent,
    }


@pytest.fixture
def applied_control_projection_world(relationship_world, monkeypatch):
    """Visible and deliberately split-authority relations for control projections."""

    world = relationship_world
    monkeypatch.setattr(FindingsAssessment, "upsert_daily_metrics", lambda self: None)
    monkeypatch.setattr(RiskAssessment, "upsert_daily_metrics", lambda self: None)

    hidden_folder = Folder.objects.create(
        name=f"hidden-control-projection-{uuid.uuid4().hex[:8]}",
        content_type=Folder.ContentType.DOMAIN,
        parent_folder=Folder.get_root_folder(),
    )
    visible_findings_assessment = FindingsAssessment.objects.create(
        name=f"visible-findings-assessment-{uuid.uuid4().hex}",
        folder=world["folder"],
    )
    hidden_findings_assessment = FindingsAssessment.objects.create(
        name=f"hidden-findings-assessment-{uuid.uuid4().hex}",
        folder=hidden_folder,
    )
    visible_finding = Finding.objects.create(
        name=f"visible-finding-{uuid.uuid4().hex}",
        folder=world["folder"],
        findings_assessment=visible_findings_assessment,
    )
    hidden_finding = Finding.objects.create(
        name=f"hidden-finding-{uuid.uuid4().hex}",
        folder=hidden_folder,
        findings_assessment=hidden_findings_assessment,
    )

    def risk_assessment(folder: Folder, visibility: str):
        matrix = RiskMatrix.objects.create(
            name=f"{visibility}-projection-matrix-{uuid.uuid4().hex}",
            folder=folder,
            json_definition={},
        )
        return RiskAssessment.objects.create(
            name=f"{visibility}-projection-risk-{uuid.uuid4().hex}",
            folder=folder,
            perimeter=Perimeter.objects.create(
                name=f"{visibility}-projection-perimeter-{uuid.uuid4().hex}",
                folder=folder,
            ),
            risk_matrix=matrix,
        )

    visible_risk_assessment = risk_assessment(world["folder"], "visible")
    hidden_risk_assessment = risk_assessment(hidden_folder, "hidden-parent")
    visible_scenario = RiskScenario.objects.create(
        name=f"visible-projection-scenario-{uuid.uuid4().hex}",
        ref_id=f"VISIBLE-{uuid.uuid4().hex[:8]}",
        folder=world["folder"],
        risk_assessment=visible_risk_assessment,
    )
    parent_hidden_scenario = RiskScenario.objects.create(
        name=f"parent-hidden-projection-scenario-{uuid.uuid4().hex}",
        ref_id=f"PARENT-HIDDEN-{uuid.uuid4().hex[:8]}",
        folder=hidden_folder,
        risk_assessment=hidden_risk_assessment,
    )
    RiskScenario.objects.filter(id=visible_scenario.id).update(
        current_level=3,
        residual_level=1,
    )
    RiskScenario.objects.filter(id=parent_hidden_scenario.id).update(
        current_level=4,
        residual_level=0,
    )
    visible_scenario.refresh_from_db()
    parent_hidden_scenario.refresh_from_db()

    # The child itself is readable in the hidden folder, but its parent risk
    # assessment is not. The projection must require both authorities.
    _grant_permissions(
        world["respondent"],
        hidden_folder,
        "view_riskscenario",
    )

    return {
        **world,
        "hidden_folder": hidden_folder,
        "visible_finding": visible_finding,
        "hidden_finding": hidden_finding,
        "visible_risk_assessment": visible_risk_assessment,
        "hidden_risk_assessment": hidden_risk_assessment,
        "visible_scenario": visible_scenario,
        "parent_hidden_scenario": parent_hidden_scenario,
    }


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
def test_inverse_relationship_rejects_forged_unassigned_addition(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    serializer = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_a"], world["unassigned"]),
    )

    with pytest.raises(PermissionDenied):
        serializer.is_valid(raise_exception=True)

    assert set(target.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id
    }


@pytest.mark.parametrize("policy_field,target_model,serializer_class", TARGETS)
def test_inverse_relationship_preserves_hidden_existing_and_projects_response(
    relationship_world,
    policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    # The unassigned row is generic-IAM visible but outside respondent actor
    # scope.  It must survive an update without appearing in the response.
    target.requirement_assessments.add(
        world["assigned_a"],
        world["unassigned"],
    )
    serializer = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
    )

    serializer.is_valid(raise_exception=True)
    saved = serializer.save()

    assert set(saved.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_b"].id,
        world["unassigned"].id,
    }
    assert {str(value) for value in serializer.data["requirement_assessments"]} == {
        str(world["assigned_b"].id)
    }
    assert str(world["unassigned"].id) not in {
        str(value) for value in serializer.data["requirement_assessments"]
    }
    read_serializer = READ_SERIALIZERS[target_model](
        saved,
        context={"request": _request_for(world["respondent"])},
    )
    assert _represented_ids(read_serializer.data["requirement_assessments"]) == {
        str(world["assigned_b"].id)
    }
    if target_model is SecurityException:
        assert read_serializer.data["associated_objects_count"] == 1
    assert policy_field in world["assessment"].field_visibility


@pytest.mark.parametrize("policy_field,target_model,serializer_class", TARGETS)
@pytest.mark.parametrize(
    "denial",
    ("read_only", "locked", "in_review", "terminal_assignment"),
)
def test_inverse_relationship_denies_non_editable_or_closed_delta(
    relationship_world,
    policy_field,
    target_model,
    serializer_class,
    denial,
):
    world = relationship_world
    assessment = world["assessment"]
    if denial == "read_only":
        assessment.field_visibility = {
            **assessment.field_visibility,
            policy_field: {"auditor": "edit", "respondent": "read"},
        }
        assessment.save(update_fields=["field_visibility"])
    elif denial == "locked":
        assessment.is_locked = True
        assessment.save(update_fields=["is_locked"])
    elif denial == "in_review":
        assessment.status = Assessment.Status.IN_REVIEW
        assessment.save(update_fields=["status"])
    else:
        world["assignment"].status = RequirementAssignment.Status.SUBMITTED
        world["assignment"].save(update_fields=["status"])

    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    serializer = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
    )

    serializer.is_valid(raise_exception=True)
    with pytest.raises(PermissionDenied):
        serializer.save()

    assert set(target.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id
    }


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
def test_inverse_relationship_requires_authenticated_request(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    serializer = _serializer(
        serializer_class,
        target=target,
        user=None,
        requested=(world["assigned_a"],),
    )

    with pytest.raises(NotAuthenticated):
        serializer.is_valid(raise_exception=True)


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
def test_inverse_relationship_snapshot_change_fails_closed(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    serializer = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
    )
    serializer.is_valid(raise_exception=True)

    # Simulate a committed competing writer between validation and save.
    target.requirement_assessments.add(world["unassigned"])
    with pytest.raises(RequirementAssessmentRelationshipConflict):
        serializer.save()

    assert set(target.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id,
        world["unassigned"].id,
    }


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
@pytest.mark.parametrize("competing_change", ("terminal_status", "actor_unlink"))
def test_inverse_relationship_assignment_authority_change_fails_closed(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
    competing_change,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    serializer = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
    )
    serializer.is_valid(raise_exception=True)

    # These direct writes simulate a competing transaction that commits after
    # validation.  PostgreSQL acceptance separately exercises actual blocking;
    # this deterministic test proves the lock-time snapshot comparison.
    if competing_change == "terminal_status":
        world["assignment"].status = RequirementAssignment.Status.SUBMITTED
        world["assignment"].save(update_fields=["status"])
    else:
        world["assignment"].actor.clear()

    with pytest.raises(RequirementAssessmentRelationshipConflict):
        serializer.save()

    assert set(target.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id
    }


def test_target_change_permission_is_reproved_before_commit(
    relationship_world,
    monkeypatch,
):
    world = relationship_world
    target = _target(Evidence, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    serializer = _serializer(
        EvidenceWriteSerializer,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
    )
    serializer.is_valid(raise_exception=True)
    from core import requirement_assessment_relationships

    original_target_action_allowed = (
        requirement_assessment_relationships._target_action_allowed
    )
    target_checks = 0

    def changing_target_authority(*, user, target, action):
        nonlocal target_checks
        if type(target) is Evidence and action == "change":
            target_checks += 1
            if target_checks > 1:
                return False
        return original_target_action_allowed(
            user=user,
            target=target,
            action=action,
        )

    monkeypatch.setattr(
        requirement_assessment_relationships,
        "_target_action_allowed",
        changing_target_authority,
    )

    with pytest.raises(RequirementAssessmentRelationshipConflict):
        serializer.save()

    assert target_checks == 2
    assert set(target.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id
    }


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
def test_real_serializer_factory_uses_governed_write_and_projected_read(
    _policy_field,
    target_model,
    serializer_class,
):
    factory = SerializerFactory("core.serializers")
    resolved_write = factory.get_serializer(target_model.__name__, "partial_update")
    resolved_read = factory.get_serializer(target_model.__name__, "retrieve")

    assert resolved_write is serializer_class
    assert issubclass(resolved_write, RequirementAssessmentRelationshipAuthorityMixin)
    assert resolved_read is READ_SERIALIZERS[target_model]


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
def test_governed_create_persists_target_and_relationship_atomically(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    name = f"created-inverse-target-{uuid.uuid4().hex}"
    serializer = serializer_class(
        data={
            "name": name,
            "folder": str(world["folder"].id),
            "requirement_assessments": [str(world["assigned_a"].id)],
        },
        context={"request": _request_for(world["respondent"])},
    )

    serializer.is_valid(raise_exception=True)
    created = serializer.save()

    assert target_model.objects.filter(id=created.id, name=name).exists()
    assert set(created.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id
    }


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
@pytest.mark.parametrize("relationship_value", ("hidden", "missing", "malformed"))
def test_governed_create_relationship_operand_fails_closed(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
    relationship_value,
):
    world = relationship_world
    name = f"rejected-inverse-target-{uuid.uuid4().hex}"
    values = {
        "hidden": str(world["unassigned"].id),
        "missing": str(uuid.uuid4()),
        "malformed": "not-a-uuid",
    }
    serializer = serializer_class(
        data={
            "name": name,
            "folder": str(world["folder"].id),
            "requirement_assessments": [values[relationship_value]],
        },
        context={"request": _request_for(world["respondent"])},
    )

    with pytest.raises(PermissionDenied):
        serializer.is_valid(raise_exception=True)
    assert not target_model.objects.filter(name=name).exists()


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
def test_ordinary_create_without_governed_relationship_remains_available(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    name = f"ordinary-create-{uuid.uuid4().hex}"
    serializer = serializer_class(
        data={"name": name, "folder": str(world["folder"].id)},
        context={"request": _request_for(world["respondent"])},
    )

    serializer.is_valid(raise_exception=True)
    created = serializer.save()

    assert target_model.objects.filter(id=created.id, name=name).exists()
    assert not created.requirement_assessments.exists()


def test_governed_create_rolls_back_target_when_final_reproof_fails(
    relationship_world,
    monkeypatch,
):
    world = relationship_world
    name = f"rolled-back-inverse-target-{uuid.uuid4().hex}"
    serializer = EvidenceWriteSerializer(
        data={
            "name": name,
            "folder": str(world["folder"].id),
            "requirement_assessments": [str(world["assigned_a"].id)],
        },
        context={"request": _request_for(world["respondent"])},
    )
    serializer.is_valid(raise_exception=True)

    def fail_final_reproof(_context, _plan):
        raise RequirementAssessmentRelationshipConflict()

    monkeypatch.setattr(
        "core.requirement_assessment_relationships._revalidate_authority_snapshot",
        fail_final_reproof,
    )
    with pytest.raises(RequirementAssessmentRelationshipConflict):
        serializer.save()

    assert not Evidence.objects.filter(name=name).exists()


def test_governed_create_reproves_target_add_permission_before_commit(
    relationship_world,
    monkeypatch,
):
    from core import requirement_assessment_relationships

    world = relationship_world
    name = f"revoked-create-target-{uuid.uuid4().hex}"
    serializer = EvidenceWriteSerializer(
        data={
            "name": name,
            "folder": str(world["folder"].id),
            "requirement_assessments": [str(world["assigned_a"].id)],
        },
        context={"request": _request_for(world["respondent"])},
    )
    serializer.is_valid(raise_exception=True)
    original_target_action_allowed = (
        requirement_assessment_relationships._target_action_allowed
    )

    def revoked_add(*, user, target, action):
        if action == "add":
            return False
        return original_target_action_allowed(
            user=user,
            target=target,
            action=action,
        )

    monkeypatch.setattr(
        requirement_assessment_relationships,
        "_target_action_allowed",
        revoked_add,
    )

    with pytest.raises(RequirementAssessmentRelationshipConflict):
        serializer.save()
    assert not Evidence.objects.filter(name=name).exists()


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
def test_governed_update_rejects_scalar_in_raw_payload(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])
    new_name = f"mixed-update-{uuid.uuid4().hex}"
    serializer = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
        name=new_name,
    )

    assert not serializer.is_valid()
    assert list(serializer.errors) == ["requirement_assessments"]
    target.refresh_from_db()

    assert target.name != new_name
    assert set(target.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id,
        world["unassigned"].id,
    }


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
def test_ordinary_scalar_update_accepts_unchanged_projected_relationship_echo(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])
    new_name = f"ordinary-update-{uuid.uuid4().hex}"
    serializer = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_a"],),
        name=new_name,
    )

    serializer.is_valid(raise_exception=True)
    updated = serializer.save()

    assert updated.name == new_name
    assert set(updated.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id,
        world["unassigned"].id,
    }


@pytest.mark.parametrize(
    "extra_payload",
    (
        {"name": "must-not-parse"},
        {"folder": "not-a-uuid"},
        {"attachment": "not-a-file"},
    ),
)
def test_governed_update_rejects_scalar_fk_or_file_before_any_field_parsing(
    relationship_world,
    extra_payload,
):
    world = relationship_world
    target = _target(Evidence, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])

    def serializer_for(relationship_id):
        return EvidenceWriteSerializer(
            target,
            data={
                "requirement_assessments": [relationship_id],
                **extra_payload,
            },
            partial=True,
            context={"request": _request_for(world["respondent"])},
        )

    visible = serializer_for(str(world["assigned_b"].id))
    malformed = serializer_for("not-a-uuid")
    assert not visible.is_valid()
    assert not malformed.is_valid()
    assert visible.errors == malformed.errors
    assert list(visible.errors) == ["requirement_assessments"]


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
@pytest.mark.parametrize("operation", ("add", "remove"))
def test_governed_delta_applies_after_hidden_baseline_is_locked(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
    operation,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])
    operand = world["assigned_b"] if operation == "add" else world["assigned_a"]
    serializer = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(operand,),
        operation=operation,
    )

    serializer.is_valid(raise_exception=True)
    serializer.save()

    expected = {world["unassigned"].id}
    if operation == "add":
        expected.update({world["assigned_a"].id, world["assigned_b"].id})
    assert set(target.requirement_assessments.values_list("id", flat=True)) == expected


def test_governed_delta_rejects_relationship_change_after_validation(
    relationship_world,
):
    world = relationship_world
    target = _target(Evidence, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    serializer = _serializer(
        EvidenceWriteSerializer,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
        operation="add",
    )
    serializer.is_valid(raise_exception=True)

    target.requirement_assessments.add(world["unassigned"])
    with pytest.raises(RequirementAssessmentRelationshipConflict):
        serializer.save()

    assert set(target.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id,
        world["unassigned"].id,
    }


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
@pytest.mark.parametrize("action", ("add_m2m", "remove_m2m"))
def test_batch_delta_preserves_hidden_baseline_without_absolute_preread(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
    action,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])
    operand = world["assigned_b"] if action == "add_m2m" else world["assigned_a"]
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.post(
        f"/api/{ENDPOINTS[target_model]}/batch-action/",
        {
            "action": action,
            "ids": [str(target.id)],
            "field": "requirement_assessments",
            "value": [str(operand.id)],
        },
        format="json",
    )

    assert response.status_code == 200, response.content
    assert response.json()["succeeded"] == [{"id": str(target.id), "name": target.name}]
    expected = {world["unassigned"].id}
    if action == "add_m2m":
        expected.update({world["assigned_a"].id, world["assigned_b"].id})
    assert set(target.requirement_assessments.values_list("id", flat=True)) == expected


@pytest.mark.parametrize("_policy_field,target_model,_serializer_class", TARGETS)
def test_detail_api_projects_inverse_relationship_and_derived_metadata(
    relationship_world,
    _policy_field,
    target_model,
    _serializer_class,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get(f"/api/{ENDPOINTS[target_model]}/{target.id}/")

    assert response.status_code == 200, response.content
    payload = response.json()
    assert _represented_ids(payload["requirement_assessments"]) == {
        str(world["assigned_a"].id)
    }
    if target_model is SecurityException:
        assert payload["associated_objects_count"] == 1


@pytest.mark.parametrize("target_model", (AppliedControl, Policy))
def test_full_api_projects_inverse_relationships(target_model, relationship_world):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get(f"/api/{ENDPOINTS[target_model]}/full/")

    assert response.status_code == 200, response.content
    body = response.json()
    rows = body.get("results", body)
    payload = next(row for row in rows if row["id"] == str(target.id))
    assert _represented_ids(payload["requirement_assessments"]) == {
        str(world["assigned_a"].id)
    }


def test_applied_control_list_and_link_filters_use_projected_relationships(
    relationship_world,
):
    world = relationship_world
    visible_target = _target(AppliedControl, world["folder"])
    visible_target.requirement_assessments.add(world["assigned_a"])
    hidden_only_target = _target(AppliedControl, world["folder"])
    hidden_only_target.requirement_assessments.add(world["unassigned"])
    client = APIClient()
    client.force_authenticate(world["respondent"])

    list_response = client.get("/api/applied-controls/")
    linked_response = client.get(
        "/api/applied-controls/",
        {"linked_models": "requirement_assessments"},
    )
    orphan_response = client.get(
        "/api/applied-controls/",
        {"linked_models": "--"},
    )

    assert list_response.status_code == 200, list_response.content
    assert linked_response.status_code == 200, linked_response.content
    assert orphan_response.status_code == 200, orphan_response.content

    def rows(response):
        body = response.json()
        return body.get("results", body)

    listed = {row["id"]: row for row in rows(list_response)}
    assert "requirement_assessments" in listed[str(visible_target.id)]["linked_models"]
    assert (
        "requirement_assessments"
        not in listed[str(hidden_only_target.id)]["linked_models"]
    )
    linked_ids = {row["id"] for row in rows(linked_response)}
    orphan_ids = {row["id"] for row in rows(orphan_response)}
    assert str(visible_target.id) in linked_ids
    assert str(hidden_only_target.id) not in linked_ids
    assert str(visible_target.id) not in orphan_ids
    assert str(hidden_only_target.id) in orphan_ids


def test_applied_control_detail_metrics_require_related_object_and_parent_iam(
    applied_control_projection_world,
):
    world = applied_control_projection_world
    visible_target = _target(AppliedControl, world["folder"])
    visible_target.effort = "S"
    visible_target.save(update_fields=["effort"])
    hidden_only_target = _target(AppliedControl, world["folder"])
    hidden_only_target.effort = "S"
    hidden_only_target.save(update_fields=["effort"])

    world["visible_finding"].applied_controls.add(visible_target)
    world["hidden_finding"].applied_controls.add(
        visible_target,
        hidden_only_target,
    )
    world["visible_scenario"].applied_controls.add(visible_target)
    world["parent_hidden_scenario"].applied_controls.add(
        visible_target,
        hidden_only_target,
    )

    user = world["respondent"]
    assert RoleAssignment.is_object_readable(
        user,
        RiskScenario,
        world["parent_hidden_scenario"].id,
    )
    assert not RoleAssignment.is_object_readable(
        user,
        RiskAssessment,
        world["hidden_risk_assessment"].id,
    )
    client = APIClient()
    client.force_authenticate(user)

    visible_response = client.get(f"/api/applied-controls/{visible_target.id}/")
    hidden_response = client.get(f"/api/applied-controls/{hidden_only_target.id}/")

    assert visible_response.status_code == 200, visible_response.content
    assert hidden_response.status_code == 200, hidden_response.content
    visible_payload = visible_response.json()
    hidden_payload = hidden_response.json()
    assert _represented_ids(visible_payload["findings"]) == {
        str(world["visible_finding"].id)
    }
    assert visible_payload["findings_count"] == 1
    assert visible_payload["ranking_score"] == 6
    assert {"findings", "risk_scenarios"} <= set(visible_payload["linked_models"])
    assert hidden_payload["findings"] == []
    assert hidden_payload["findings_count"] == 0
    assert hidden_payload["ranking_score"] == 0
    assert "findings" not in hidden_payload["linked_models"]
    assert "risk_scenarios" not in hidden_payload["linked_models"]
    assert world["hidden_finding"].name not in hidden_response.content.decode()
    assert world["parent_hidden_scenario"].name not in hidden_response.content.decode()


def test_applied_control_link_filters_treat_hidden_relations_as_orphaned(
    applied_control_projection_world,
):
    world = applied_control_projection_world
    visible_target = _target(AppliedControl, world["folder"])
    hidden_only_target = _target(AppliedControl, world["folder"])
    orphan_target = _target(AppliedControl, world["folder"])
    world["visible_finding"].applied_controls.add(visible_target)
    world["visible_scenario"].applied_controls.add(visible_target)
    world["hidden_finding"].applied_controls.add(hidden_only_target)
    world["parent_hidden_scenario"].applied_controls.add(hidden_only_target)
    client = APIClient()
    client.force_authenticate(world["respondent"])

    findings_response = client.get(
        "/api/applied-controls/",
        {"linked_models": "findings"},
    )
    risks_response = client.get(
        "/api/applied-controls/",
        {"linked_models": "risk_scenarios"},
    )
    orphan_response = client.get(
        "/api/applied-controls/",
        {"linked_models": "--"},
    )

    assert findings_response.status_code == 200, findings_response.content
    assert risks_response.status_code == 200, risks_response.content
    assert orphan_response.status_code == 200, orphan_response.content

    def response_ids(response):
        body = response.json()
        return {row["id"] for row in body.get("results", body)}

    assert str(visible_target.id) in response_ids(findings_response)
    assert str(hidden_only_target.id) not in response_ids(findings_response)
    assert str(visible_target.id) in response_ids(risks_response)
    assert str(hidden_only_target.id) not in response_ids(risks_response)
    assert response_ids(orphan_response) >= {
        str(hidden_only_target.id),
        str(orphan_target.id),
    }
    assert str(visible_target.id) not in response_ids(orphan_response)


def test_applied_control_list_projection_queries_do_not_scale_per_row(
    relationship_world,
):
    world = relationship_world
    targets = [_target(AppliedControl, world["folder"]) for _index in range(5)]
    for target in targets:
        target.requirement_assessments.add(
            world["assigned_a"],
            world["unassigned"],
        )
        # Mirror the list view's prefetches so this invariant isolates the
        # caller-scoped requirement-assessment projection.
        target._prefetched_objects_cache = {
            "owner": [],
            "filtering_labels": [],
            "assets": [],
        }

    def serialize_and_count(items):
        context = {"request": _request_for(world["respondent"])}
        with CaptureQueriesContext(connection) as captured:
            payload = AppliedControlListSerializer(
                items,
                many=True,
                context=context,
            ).data
        assert all("requirement_assessments" in row["linked_models"] for row in payload)
        return len(captured)

    # Warm process- and instance-local relation caches before comparing the
    # query slope; the invariant is page size, not cold-start metadata cost.
    serialize_and_count(targets[:1])
    single_count = serialize_and_count(targets[:1])
    page_count = serialize_and_count(targets)

    assert page_count == single_count


def test_rendered_relationship_projection_queries_do_not_scale_per_row(
    relationship_world,
):
    world = relationship_world
    framework = world["assessment"].framework
    framework.field_visibility = {
        "evidences": {"auditor": "edit", "respondent": "edit"}
    }
    framework.save(update_fields=["field_visibility"])
    world["assessment"].field_visibility = {}
    world["assessment"].save(update_fields=["field_visibility"])
    targets = [_target(Evidence, world["folder"]) for _index in range(5)]
    expected_by_target = {}
    for index, target in enumerate(targets):
        visible_node = RequirementNode.objects.create(
            name=f"Visible projection row {index}",
            urn=f"{framework.urn}:projection-visible-{index}",
            ref_id=f"PV-{index}",
            framework=framework,
            folder=world["folder"],
            assessable=True,
        )
        hidden_node = RequirementNode.objects.create(
            name=f"Hidden projection row {index}",
            urn=f"{framework.urn}:projection-hidden-{index}",
            ref_id=f"PH-{index}",
            framework=framework,
            folder=world["folder"],
            assessable=True,
        )
        visible_row = RequirementAssessment.objects.create(
            compliance_assessment=world["assessment"],
            requirement=visible_node,
            folder=world["folder"],
        )
        hidden_row = RequirementAssessment.objects.create(
            compliance_assessment=world["assessment"],
            requirement=hidden_node,
            folder=world["folder"],
        )
        world["assignment"].requirement_assessments.add(visible_row)
        target.requirement_assessments.add(visible_row, hidden_row)
        expected_by_target[str(target.id)] = str(visible_row.id)

    def serialize_and_count(items):
        context = {"request": _request_for(world["respondent"])}
        with CaptureQueriesContext(connection) as captured:
            payload = _ProjectionOnlyEvidenceSerializer(
                items,
                many=True,
                context=context,
            ).data
        assert all(
            _represented_ids(row["requirement_assessments"])
            == {expected_by_target[str(row["id"])]}
            for row in payload
        )
        return len(captured)

    serialize_and_count(targets[:1])
    single_count = serialize_and_count(targets[:1])
    page_count = serialize_and_count(targets)

    assert page_count == single_count


def test_cross_assessment_assignment_link_cannot_authorize_batched_projection(
    relationship_world,
):
    world = relationship_world
    other_assessment = ComplianceAssessment.objects.create(
        name="Cross-assessment assignment audit",
        ref_id="CROSS-ASSIGNMENT-AUDIT",
        framework=world["assessment"].framework,
        folder=world["folder"],
        perimeter=world["assessment"].perimeter,
        min_score=0,
        max_score=4,
        status=Assessment.Status.IN_PROGRESS,
        field_visibility=world["assessment"].field_visibility,
    )
    other_node = RequirementNode.objects.create(
        name="Cross-assessment assignment requirement",
        urn=f"{world['assessment'].framework.urn}:cross-assignment",
        ref_id="CROSS-ASSIGNMENT",
        framework=world["assessment"].framework,
        folder=world["folder"],
        assessable=True,
    )
    other_row = RequirementAssessment.objects.create(
        compliance_assessment=other_assessment,
        requirement=other_node,
        folder=world["folder"],
    )
    # Simulate a legacy/corrupt through row: the assignment's parent audit is
    # A, while the linked RequirementAssessment belongs to audit B.
    world["assignment"].requirement_assessments.add(other_row)
    visible_target = _target(Evidence, world["folder"])
    visible_target.requirement_assessments.add(world["assigned_a"])
    corrupt_target = _target(Evidence, world["folder"])
    corrupt_target.requirement_assessments.add(other_row)

    payload = _ProjectionOnlyEvidenceSerializer(
        [visible_target, corrupt_target],
        many=True,
        context={"request": _request_for(world["respondent"])},
    ).data
    payload_by_id = {str(row["id"]): row for row in payload}

    assert _represented_ids(
        payload_by_id[str(visible_target.id)]["requirement_assessments"]
    ) == {str(world["assigned_a"].id)}
    assert payload_by_id[str(corrupt_target.id)]["requirement_assessments"] == []

    actor = Actor.objects.get(user=world["respondent"])
    snapshot = _assignment_authority_snapshot(
        actor_ids=(actor.id,),
        assessment_ids=(world["assessment"].id, other_assessment.id),
        ra_ids=(world["assigned_a"].id, other_row.id),
    )
    assert {ra_id for _assignment_id, _status, ra_id in snapshot} == {
        world["assigned_a"].id
    }


def test_relationship_read_serializers_install_batch_projection():
    serializers_to_check = (
        AppliedControlReadSerializer,
        AppliedControlBulkReadSerializer,
        AppliedControlListSerializer,
        PolicyReadSerializer,
        EvidenceReadSerializer,
        SecurityExceptionReadSerializer,
    )

    for serializer_class in serializers_to_check:
        serializer = serializer_class([], many=True)
        assert isinstance(
            serializer,
            RequirementAssessmentRelationshipProjectionListSerializer,
        )


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
@pytest.mark.parametrize("include_relationship", (False, True))
def test_many_write_rejects_before_any_child_save(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
    include_relationship,
):
    world = relationship_world
    names = [f"many-write-{uuid.uuid4().hex}" for _index in range(2)]
    payload = [
        {
            "name": name,
            "folder": str(world["folder"].id),
        }
        for name in names
    ]
    if include_relationship:
        for item in payload:
            item["requirement_assessments"] = [str(world["assigned_a"].id)]
    serializer = serializer_class(
        data=payload,
        many=True,
        context={"request": _request_for(world["respondent"])},
    )

    with pytest.raises(ValidationError):
        serializer.is_valid(raise_exception=True)

    assert not target_model.objects.filter(name__in=names).exists()


def test_object_action_projects_relationship_with_request_context(relationship_world):
    world = relationship_world
    target = _target(Evidence, world["folder"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get(f"/api/evidences/{target.id}/object/")

    assert response.status_code == 200, response.content
    assert _represented_ids(response.json()["requirement_assessments"]) == {
        str(world["assigned_a"].id)
    }


def test_todo_action_passes_request_to_applied_control_projection(
    relationship_world,
):
    world = relationship_world
    target = _target(AppliedControl, world["folder"])
    target.eta = timezone.localdate()
    target.save(update_fields=["eta"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get("/api/applied-controls/todo/")

    assert response.status_code == 200, response.content
    payload = next(
        item for item in response.json()["results"] if item["id"] == str(target.id)
    )
    assert _represented_ids(payload["requirement_assessments"]) == {
        str(world["assigned_a"].id)
    }
    assert "requirement_assessments" in payload["linked_models"]


def test_todo_action_masks_related_objects_outside_caller_iam(
    relationship_world,
):
    world = relationship_world
    hidden_folder = Folder.objects.create(
        name=f"hidden-action-relation-{uuid.uuid4().hex[:8]}",
        content_type=Folder.ContentType.DOMAIN,
        parent_folder=Folder.get_root_folder(),
    )
    hidden_evidence = Evidence.objects.create(
        name=f"hidden-action-evidence-{uuid.uuid4().hex}",
        folder=hidden_folder,
    )
    target = _target(AppliedControl, world["folder"])
    target.eta = timezone.localdate()
    target.save(update_fields=["eta"])
    target.evidences.add(hidden_evidence)
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get("/api/applied-controls/todo/")

    assert response.status_code == 200, response.content
    payload = next(
        item for item in response.json()["results"] if item["id"] == str(target.id)
    )
    serialized_evidences = str(payload["evidences"])
    assert str(hidden_evidence.id) not in serialized_evidences
    assert hidden_evidence.name not in serialized_evidences


def test_todo_action_preserves_non_requirement_link_annotations(
    relationship_world,
):
    world = relationship_world
    target = _target(AppliedControl, world["folder"])
    target.eta = timezone.localdate()
    target.save(update_fields=["eta"])
    task_template = TaskTemplate.objects.create(
        name=f"linked-action-task-{uuid.uuid4().hex}",
        folder=world["folder"],
    )
    task_template.applied_controls.add(target)
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get("/api/applied-controls/todo/")

    assert response.status_code == 200, response.content
    payload = next(
        item for item in response.json()["results"] if item["id"] == str(target.id)
    )
    assert "task_templates" in payload["linked_models"]


def test_priority_chart_counts_only_caller_visible_requirement_and_risk_links(
    applied_control_projection_world,
):
    world = applied_control_projection_world
    visible_target = _target(AppliedControl, world["folder"])
    visible_target.priority = 1
    visible_target.save(update_fields=["priority"])
    visible_target.requirement_assessments.add(world["assigned_a"])
    world["visible_scenario"].applied_controls.add(visible_target)

    hidden_only_target = _target(AppliedControl, world["folder"])
    hidden_only_target.priority = 1
    hidden_only_target.save(update_fields=["priority"])
    hidden_only_target.requirement_assessments.add(world["unassigned"])
    world["parent_hidden_scenario"].applied_controls.add(hidden_only_target)
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get("/api/applied-controls/priority_chart_data/")

    assert response.status_code == 200, response.content
    vectors = {
        row[5]: row
        for rows in (
            response.json()[status]
            for status in ("--", "to_do", "in_progress", "on_hold", "deprecated")
        )
        for row in rows
    }
    assert vectors[str(visible_target.id)][2] == 7
    assert vectors[str(hidden_only_target.id)][2] == 5


def test_todo_scores_and_orders_controls_from_caller_visible_risks_only(
    applied_control_projection_world,
):
    world = applied_control_projection_world
    visible_target = _target(AppliedControl, world["folder"])
    visible_target.effort = "S"
    visible_target.eta = timezone.localdate()
    visible_target.save(update_fields=["effort", "eta"])
    world["visible_scenario"].applied_controls.add(visible_target)

    hidden_only_target = _target(AppliedControl, world["folder"])
    hidden_only_target.effort = "XS"
    hidden_only_target.eta = timezone.localdate()
    hidden_only_target.save(update_fields=["effort", "eta"])
    world["parent_hidden_scenario"].applied_controls.add(hidden_only_target)
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get("/api/applied-controls/todo/")

    assert response.status_code == 200, response.content
    rows = response.json()["results"]
    rows_by_id = {row["id"]: row for row in rows}
    returned_ids = [row["id"] for row in rows]
    assert rows_by_id[str(visible_target.id)]["ranking_score"] == 6
    assert rows_by_id[str(hidden_only_target.id)]["ranking_score"] == 0
    assert returned_ids.index(str(visible_target.id)) < returned_ids.index(
        str(hidden_only_target.id)
    )


def test_risk_action_plan_excludes_scenario_with_hidden_parent(
    applied_control_projection_world,
):
    world = applied_control_projection_world
    target = _target(AppliedControl, world["folder"])
    target.effort = "S"
    target.save(update_fields=["effort"])
    hidden_only_target = _target(AppliedControl, world["folder"])
    world["visible_scenario"].applied_controls.add(target)
    world["parent_hidden_scenario"].applied_controls.add(
        target,
        hidden_only_target,
    )
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get(
        f"/api/risk-assessments/{world['visible_risk_assessment'].id}/action-plan/"
    )

    assert response.status_code == 200, response.content
    body = response.json()
    rows = body.get("results", body)
    rows_by_id = {row["id"]: row for row in rows}
    assert str(target.id) in rows_by_id
    assert str(hidden_only_target.id) not in rows_by_id
    assert rows_by_id[str(target.id)]["risk_scenarios"] == [
        {
            "str": (
                f"{world['visible_scenario'].ref_id} - {world['visible_scenario'].name}"
            ),
            "id": str(world["visible_scenario"].id),
        }
    ]
    assert rows_by_id[str(target.id)]["ranking_score"] == 6
    rendered = response.content.decode()
    assert world["parent_hidden_scenario"].name not in rendered
    assert world["parent_hidden_scenario"].ref_id not in rendered


def test_risk_action_plan_installs_batch_request_projection():
    serializer = RiskAssessmentActionPlanSerializer([], many=True)

    assert isinstance(serializer, AppliedControlRequestProjectionListSerializer)


def test_impact_graph_projects_requirement_and_risk_parent_iam(
    relationship_world,
):
    world = relationship_world
    target = _target(AppliedControl, world["folder"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])

    hidden_folder = Folder.objects.create(
        name=f"hidden-impact-graph-{uuid.uuid4().hex[:8]}",
        content_type=Folder.ContentType.DOMAIN,
        parent_folder=Folder.get_root_folder(),
    )
    risk_matrix = RiskMatrix.objects.create(
        name=f"hidden-matrix-{uuid.uuid4().hex}",
        folder=hidden_folder,
        json_definition={},
    )
    risk_assessment = RiskAssessment.objects.create(
        name=f"hidden-risk-assessment-{uuid.uuid4().hex}",
        folder=hidden_folder,
        perimeter=Perimeter.objects.create(
            name=f"hidden-risk-perimeter-{uuid.uuid4().hex}",
            folder=hidden_folder,
        ),
        risk_matrix=risk_matrix,
    )
    scenario = RiskScenario.objects.create(
        name=f"hidden-risk-scenario-{uuid.uuid4().hex}",
        ref_id=f"HIDDEN-{uuid.uuid4().hex[:8]}",
        folder=hidden_folder,
        risk_assessment=risk_assessment,
    )
    scenario.applied_controls.add(target)
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get("/api/applied-controls/impact_graph/")

    assert response.status_code == 200, response.content
    content = response.content.decode()
    assert world["assigned_a"].requirement.ref_id in content
    assert world["unassigned"].requirement.ref_id not in content
    assert scenario.name not in content
    assert scenario.ref_id not in content
    assert risk_assessment.name not in content


@pytest.mark.parametrize("action_path", ("todo", "to_review"))
def test_unpaginated_applied_control_action_queries_do_not_scale_per_row(
    relationship_world,
    action_path,
    monkeypatch,
):
    world = relationship_world
    client = APIClient()
    client.force_authenticate(world["respondent"])
    from core import serializers as core_serializers

    original_ff_is_enabled = core_serializers.ff_is_enabled
    monkeypatch.setattr(
        core_serializers,
        "ff_is_enabled",
        lambda name: True if name == "custom_fields" else original_ff_is_enabled(name),
    )
    content_type = ContentType.objects.get_for_model(AppliedControl)
    definition = CustomFieldDefinition.objects.create(
        key=f"action-slope-{uuid.uuid4().hex[:8]}",
        label="Action slope",
        field_type=FieldType.TEXT,
        content_type=content_type,
        folder=world["folder"],
    )

    def create_target(index):
        child_folder = Folder.objects.create(
            name=f"action-slope-folder-{index}-{uuid.uuid4().hex[:8]}",
            content_type=Folder.ContentType.DOMAIN,
            parent_folder=world["folder"],
        )
        target = _target(AppliedControl, child_folder)
        target.eta = timezone.localdate()
        target.expiry_date = timezone.localdate()
        target.save(update_fields=["eta", "expiry_date"])
        target.requirement_assessments.add(world["assigned_a"])
        CustomFieldValue.objects.create(
            definition=definition,
            content_type=content_type,
            object_id=target.id,
            value_text=f"value-{index}",
        )
        return target

    targets = [create_target(0)]

    def request_and_count():
        with CaptureQueriesContext(connection) as captured:
            response = client.get(f"/api/applied-controls/{action_path}/")
        assert response.status_code == 200, response.content
        returned_rows = {row["id"]: row for row in response.json()["results"]}
        returned_ids = set(returned_rows)
        assert {str(target.id) for target in targets} <= returned_ids
        for index, target in enumerate(targets):
            assert returned_rows[str(target.id)]["custom_fields"] == {
                definition.key: f"value-{index}"
            }
        return len(captured)

    request_and_count()
    single_count = request_and_count()
    targets.extend(create_target(index) for index in range(1, 5))
    page_count = request_and_count()

    assert page_count == single_count


def test_contextless_applied_control_projection_is_explicitly_fail_closed(
    applied_control_projection_world,
):
    world = applied_control_projection_world
    target = _target(AppliedControl, world["folder"])
    target.effort = "S"
    target.save(update_fields=["effort"])
    target.requirement_assessments.add(world["assigned_a"])
    world["visible_finding"].applied_controls.add(target)
    world["visible_scenario"].applied_controls.add(target)

    payload = AppliedControlReadSerializer(target).data
    action_plan_payload = RiskAssessmentActionPlanSerializer(
        target,
        context={"pk": world["visible_risk_assessment"].id},
    ).data

    assert payload["requirement_assessments"] == []
    assert payload["findings_count"] == 0
    assert payload["ranking_score"] == 0
    assert "requirement_assessments" not in payload["linked_models"]
    assert "findings" not in payload["linked_models"]
    assert "risk_scenarios" not in payload["linked_models"]
    assert action_plan_payload["ranking_score"] == 0
    assert action_plan_payload["risk_scenarios"] == []


def test_applied_control_permission_does_not_authorize_policy_proxy(
    relationship_world,
):
    world = relationship_world
    policy = _target(Policy, world["folder"])
    applied_only_user = User.objects.create_user(
        email=f"applied-only-{uuid.uuid4().hex}@tests.invalid"
    )
    role = Role.objects.create(
        name=f"Applied only {uuid.uuid4().hex}",
        folder=Folder.get_root_folder(),
    )
    role.permissions.set(Permission.objects.filter(codename="view_appliedcontrol"))
    assignment = RoleAssignment.objects.create(
        user=applied_only_user,
        role=role,
        folder=Folder.get_root_folder(),
        is_recursive=True,
    )
    assignment.perimeter_folders.add(world["folder"])
    client = APIClient()
    client.force_authenticate(applied_only_user)

    response = client.get(f"/api/policies/{policy.id}/")

    assert response.status_code in {403, 404}
    assert str(policy.id) not in response.content.decode()


def test_applied_control_route_cannot_read_mutate_or_create_policy_proxy(
    relationship_world,
):
    world = relationship_world
    policy = _target(Policy, world["folder"])
    policy.eta = timezone.localdate()
    policy.expiry_date = timezone.localdate()
    policy.save(update_fields=["eta", "expiry_date"])
    policy.requirement_assessments.add(world["assigned_a"])
    source = _target(AppliedControl, world["folder"])
    source.eta = timezone.localdate()
    source.expiry_date = timezone.localdate()
    source.save(update_fields=["eta", "expiry_date"])
    respondent_assignment = RoleAssignment.objects.get(user=world["respondent"])
    respondent_assignment.role.permissions.remove(
        *Permission.objects.filter(
            codename__in={"view_policy", "add_policy", "change_policy"}
        )
    )
    client = APIClient()
    client.force_authenticate(world["respondent"])

    detail = client.get(f"/api/applied-controls/{policy.id}/")
    listing = client.get("/api/applied-controls/")
    mutation = client.patch(
        f"/api/applied-controls/{policy.id}/",
        {"requirement_assessments": [str(world["assigned_b"].id)]},
        format="json",
    )
    created_name = f"forged-policy-{uuid.uuid4().hex}"
    create = client.post(
        "/api/applied-controls/",
        {
            "name": created_name,
            "folder": str(world["folder"].id),
            "category": "policy",
        },
        format="json",
    )
    merge_name = f"forged-merged-policy-{uuid.uuid4().hex}"
    merge = client.post(
        "/api/applied-controls/merge/",
        {
            "source_ids": [str(source.id)],
            "target": {
                "type": "new",
                "fields": {
                    "name": merge_name,
                    "folder": str(world["folder"].id),
                    "category": "policy",
                },
            },
        },
        format="json",
    )
    policy_todo = client.get("/api/policies/todo/")
    policy_review = client.get("/api/policies/to_review/")
    policy_updatables = client.get("/api/policies/updatables/")
    concrete_updatables = client.get("/api/applied-controls/updatables/")

    listed_rows = listing.json().get("results", listing.json())
    assert detail.status_code == mutation.status_code == 404
    assert listing.status_code == 200
    assert str(policy.id) not in {row["id"] for row in listed_rows}
    assert create.status_code == merge.status_code == 400
    assert policy_todo.status_code == policy_review.status_code == 200
    assert policy_todo.json()["results"] == []
    assert policy_review.json()["results"] == []
    assert policy_updatables.status_code == 200
    assert policy_updatables.json()["results"] == []
    assert concrete_updatables.status_code == 200
    assert str(policy.id) not in {
        str(item_id) for item_id in concrete_updatables.json()["results"]
    }
    assert not AppliedControl.objects.filter(
        name__in=(created_name, merge_name),
        category="policy",
    ).exists()
    assert set(policy.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id
    }


@pytest.mark.parametrize(
    "action_path",
    (
        "get_controls_info",
        "priority_chart_data",
        "get_gantt_data",
        "impact_effort",
        "get_timeline_info",
        "ids",
        "impact_graph",
        "export_csv",
        "sunburst_data",
    ),
)
def test_policy_collection_actions_do_not_use_applied_control_authority(
    relationship_world,
    action_path,
):
    world = relationship_world
    control = _target(AppliedControl, world["folder"])
    control.start_date = timezone.localdate()
    control.eta = timezone.localdate()
    control.priority = 2
    control.control_impact = 3
    control.effort = "M"
    control.status = AppliedControl.Status.TO_DO
    control.csf_function = "identify"
    control.save(
        update_fields=[
            "start_date",
            "eta",
            "priority",
            "control_impact",
            "effort",
            "status",
            "csf_function",
        ]
    )
    respondent_role = RoleAssignment.objects.get(user=world["respondent"]).role
    respondent_role.permissions.remove(
        *Permission.objects.filter(
            codename__in={
                "view_policy",
                "add_policy",
                "change_policy",
                "delete_policy",
            }
        )
    )
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get(f"/api/policies/{action_path}/")

    assert response.status_code == 200, (action_path, response.content)
    if action_path == "sunburst_data":
        assert response.json()["results"] == []
    else:
        content = response.content.decode(errors="ignore")
        assert str(control.id) not in content
        assert control.name not in content


@pytest.mark.parametrize("action_path", ("export_csv", "export_xlsx", "mss_xlsx"))
def test_applied_control_exports_exclude_policy_proxy_rows(
    relationship_world,
    action_path,
):
    world = relationship_world
    control = _target(AppliedControl, world["folder"])
    policy = _target(Policy, world["folder"])
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get(f"/api/applied-controls/{action_path}/")

    assert response.status_code == 200, response.content
    if action_path == "export_csv":
        rendered = response.content.decode(errors="ignore")
    else:
        workbook = load_workbook(io.BytesIO(response.content), read_only=True)
        rendered = "\n".join(
            str(cell)
            for worksheet in workbook.worksheets
            for row in worksheet.iter_rows(values_only=True)
            for cell in row
            if cell is not None
        )
    assert control.name in rendered
    assert policy.name not in rendered


def test_policy_merge_requires_policy_permissions_for_existing_target(
    relationship_world,
):
    world = relationship_world
    source = _target(Policy, world["folder"])
    target = _target(Policy, world["folder"])
    role = RoleAssignment.objects.get(user=world["respondent"]).role
    role.permissions.add(
        *Permission.objects.filter(
            codename__in={
                "add_appliedcontrol",
                "change_appliedcontrol",
                "delete_appliedcontrol",
            }
        )
    )
    role.permissions.remove(
        *Permission.objects.filter(
            codename__in={"add_policy", "change_policy", "delete_policy"}
        )
    )
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.post(
        "/api/policies/merge/",
        {
            "source_ids": [str(source.id)],
            "target": {"type": "existing", "id": str(target.id)},
        },
        format="json",
    )

    assert response.status_code == 403, response.content
    assert Policy.objects.filter(id=source.id).exists()
    assert Policy.objects.filter(id=target.id).exists()


def test_policy_merge_uses_policy_serializer_for_new_target(
    relationship_world,
):
    world = relationship_world
    source = _target(Policy, world["folder"])
    role = RoleAssignment.objects.get(user=world["respondent"]).role
    role.permissions.add(
        *Permission.objects.filter(
            codename__in={"add_policy", "change_policy", "delete_policy"}
        )
    )
    role.permissions.remove(
        *Permission.objects.filter(
            codename__in={
                "add_appliedcontrol",
                "change_appliedcontrol",
                "delete_appliedcontrol",
            }
        )
    )
    client = APIClient()
    client.force_authenticate(world["respondent"])
    target_name = f"merged-policy-{uuid.uuid4().hex}"

    response = client.post(
        "/api/policies/merge/",
        {
            "source_ids": [str(source.id)],
            "target": {
                "type": "new",
                "fields": {
                    "name": target_name,
                    "folder": str(world["folder"].id),
                },
            },
        },
        format="json",
    )

    assert response.status_code == 200, response.content
    target = Policy.objects.get(id=response.json()["target_id"])
    assert target.name == target_name
    assert target.category == "policy"
    assert not Policy.objects.filter(id=source.id).exists()


@pytest.mark.parametrize(
    "target_model,endpoint",
    ((AppliedControl, "applied-controls"), (Policy, "policies")),
)
def test_duplicate_requires_exact_model_add_permission_on_target_folder(
    relationship_world,
    target_model,
    endpoint,
):
    world = relationship_world
    source = _target(target_model, world["folder"])
    hidden_folder = Folder.objects.create(
        name=f"duplicate-hidden-{uuid.uuid4().hex[:8]}",
        content_type=Folder.ContentType.DOMAIN,
        parent_folder=Folder.get_root_folder(),
    )
    client = APIClient()
    client.force_authenticate(world["respondent"])
    duplicate_name = f"unauthorized-duplicate-{uuid.uuid4().hex}"

    response = client.post(
        f"/api/{endpoint}/{source.id}/duplicate/",
        {
            "name": duplicate_name,
            "description": "",
            "folder": str(hidden_folder.id),
            "duplicate_evidences": False,
        },
        format="json",
    )

    assert response.status_code == 403, response.content
    assert not target_model.objects.filter(name=duplicate_name).exists()


@pytest.mark.parametrize(
    "target_model,endpoint,change_codename",
    (
        (AppliedControl, "applied-controls", "change_appliedcontrol"),
        (Policy, "policies", "change_policy"),
    ),
)
def test_sync_to_reference_control_requires_exact_model_change_permission(
    relationship_world,
    target_model,
    endpoint,
    change_codename,
):
    world = relationship_world
    reference = ReferenceControl.objects.create(
        name=f"sync-reference-{uuid.uuid4().hex}",
        urn=f"urn:test:sync-reference:{uuid.uuid4().hex}",
        folder=world["folder"],
        csf_function="identify",
        category="technical",
    )
    target = _target(target_model, world["folder"])
    target.reference_control = reference
    target.csf_function = "protect"
    target.save(update_fields=["reference_control", "csf_function"])
    role = RoleAssignment.objects.get(user=world["respondent"]).role
    role.permissions.remove(*Permission.objects.filter(codename=change_codename))
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.post(
        f"/api/{endpoint}/{target.id}/sync-to-reference-control/",
        {"dry_run": False},
        format="json",
    )

    assert response.status_code == 403, response.content
    target.refresh_from_db()
    assert target.csf_function == "protect"


def test_policy_sync_to_reference_control_preserves_proxy_discriminator(
    relationship_world,
):
    world = relationship_world
    reference = ReferenceControl.objects.create(
        name=f"policy-sync-reference-{uuid.uuid4().hex}",
        urn=f"urn:test:policy-sync-reference:{uuid.uuid4().hex}",
        folder=world["folder"],
        csf_function="identify",
        category="technical",
    )
    policy = _target(Policy, world["folder"])
    policy.reference_control = reference
    policy.csf_function = "protect"
    policy.save(update_fields=["reference_control", "csf_function"])
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.post(
        f"/api/policies/{policy.id}/sync-to-reference-control/?dry_run=false",
        {},
        format="json",
    )

    assert response.status_code == 200, response.content
    policy.refresh_from_db()
    assert policy.category == "policy"
    assert policy.csf_function == reference.csf_function
    assert len(response.json()) == 1


def test_generic_control_api_rejects_policy_category_derived_from_reference(
    relationship_world,
):
    world = relationship_world
    _grant_permissions(
        world["respondent"],
        world["folder"],
        "view_referencecontrol",
    )
    policy_reference = ReferenceControl.objects.create(
        name=f"policy-reference-{uuid.uuid4().hex}",
        urn=f"urn:test:policy-reference:{uuid.uuid4().hex}",
        folder=world["folder"],
        category="policy",
        csf_function="govern",
    )
    client = APIClient()
    client.force_authenticate(world["respondent"])
    derived_name = f"derived-policy-{uuid.uuid4().hex}"

    create_response = client.post(
        "/api/applied-controls/",
        {
            "name": derived_name,
            "folder": str(world["folder"].id),
            "reference_control": str(policy_reference.id),
        },
        format="json",
    )

    control = AppliedControl.objects.create(
        name=f"ordinary-control-{uuid.uuid4().hex}",
        folder=world["folder"],
        category="technical",
        csf_function="protect",
        reference_control=policy_reference,
    )
    clear_category_response = client.patch(
        f"/api/applied-controls/{control.id}/",
        {"category": None},
        format="json",
    )
    dry_run_sync_response = client.post(
        f"/api/applied-controls/{control.id}/sync-to-reference-control/",
        {},
        format="json",
    )
    sync_response = client.post(
        f"/api/applied-controls/{control.id}/sync-to-reference-control/?dry_run=false",
        {},
        format="json",
    )

    assert create_response.status_code == 400, create_response.content
    assert clear_category_response.status_code == 400, clear_category_response.content
    assert dry_run_sync_response.status_code == 400, dry_run_sync_response.content
    assert sync_response.status_code == 400, sync_response.content
    assert not AppliedControl.objects.filter(name=derived_name).exists()
    control.refresh_from_db()
    assert control.category == "technical"
    assert control.csf_function == "protect"


def test_reference_control_bulk_sync_uses_exact_proxy_permissions(
    relationship_world,
):
    world = relationship_world
    reference = ReferenceControl.objects.create(
        name=f"bulk-sync-reference-{uuid.uuid4().hex}",
        urn=f"urn:test:bulk-sync-reference:{uuid.uuid4().hex}",
        folder=world["folder"],
        category="technical",
        csf_function="identify",
    )
    ordinary_control = AppliedControl.objects.create(
        name=f"bulk-sync-control-{uuid.uuid4().hex}",
        folder=world["folder"],
        category="process",
        csf_function="protect",
        reference_control=reference,
    )
    policy = Policy.objects.create(
        name=f"bulk-sync-policy-{uuid.uuid4().hex}",
        folder=world["folder"],
        csf_function="protect",
        reference_control=reference,
    )

    applied_only_user = User.objects.create_user(
        email=f"bulk-applied-only-{uuid.uuid4().hex}@tests.invalid"
    )
    _grant_permissions(
        applied_only_user,
        world["folder"],
        "view_referencecontrol",
        "add_referencecontrol",
        "view_appliedcontrol",
        "change_appliedcontrol",
    )
    applied_client = APIClient()
    applied_client.force_authenticate(applied_only_user)
    list_url = f"/api/reference-controls/{reference.id}/syncable-applied-controls/"
    sync_url = f"/api/reference-controls/{reference.id}/sync-applied-controls/"

    applied_preview = applied_client.get(list_url)
    applied_sync = applied_client.post(sync_url, {}, format="json")

    assert applied_preview.status_code == 200, applied_preview.content
    assert applied_sync.status_code == 200, applied_sync.content
    assert {row["id"] for row in applied_preview.json()} == {str(ordinary_control.id)}
    assert {row["id"] for row in applied_sync.json()} == {str(ordinary_control.id)}
    ordinary_control.refresh_from_db()
    policy.refresh_from_db()
    assert ordinary_control.category == reference.category
    assert ordinary_control.csf_function == reference.csf_function
    assert policy.category == "policy"
    assert policy.csf_function == "protect"

    policy_only_user = User.objects.create_user(
        email=f"bulk-policy-only-{uuid.uuid4().hex}@tests.invalid"
    )
    _grant_permissions(
        policy_only_user,
        world["folder"],
        "view_referencecontrol",
        "add_referencecontrol",
        "view_policy",
        "change_policy",
    )
    policy_client = APIClient()
    policy_client.force_authenticate(policy_only_user)

    policy_preview = policy_client.get(list_url)
    policy_sync = policy_client.post(sync_url, {}, format="json")

    assert policy_preview.status_code == 200, policy_preview.content
    assert policy_sync.status_code == 200, policy_sync.content
    assert {row["id"] for row in policy_preview.json()} == {str(policy.id)}
    assert {row["id"] for row in policy_sync.json()} == {str(policy.id)}
    policy.refresh_from_db()
    assert policy.category == "policy"
    assert policy.csf_function == reference.csf_function

    change_only_control = AppliedControl.objects.create(
        name=f"bulk-change-only-control-{uuid.uuid4().hex}",
        folder=world["folder"],
        category="process",
        csf_function="protect",
        reference_control=reference,
    )
    change_only_user = User.objects.create_user(
        email=f"bulk-change-only-{uuid.uuid4().hex}@tests.invalid"
    )
    _grant_permissions(
        change_only_user,
        world["folder"],
        "view_referencecontrol",
        "add_referencecontrol",
        "change_appliedcontrol",
    )
    change_only_client = APIClient()
    change_only_client.force_authenticate(change_only_user)

    change_only_preview = change_only_client.get(list_url)
    change_only_sync = change_only_client.post(sync_url, {}, format="json")

    assert change_only_preview.status_code == 200, change_only_preview.content
    assert change_only_sync.status_code == 200, change_only_sync.content
    assert change_only_preview.json() == []
    assert change_only_sync.json() == []
    change_only_control.refresh_from_db()
    assert change_only_control.category == "process"
    assert change_only_control.csf_function == "protect"


def test_reference_control_bulk_sync_does_not_promote_generic_control_to_policy(
    relationship_world,
):
    world = relationship_world
    policy_reference = ReferenceControl.objects.create(
        name=f"bulk-policy-reference-{uuid.uuid4().hex}",
        urn=f"urn:test:bulk-policy-reference:{uuid.uuid4().hex}",
        folder=world["folder"],
        category="policy",
        csf_function="govern",
    )
    ordinary_control = AppliedControl.objects.create(
        name=f"bulk-ordinary-control-{uuid.uuid4().hex}",
        folder=world["folder"],
        category="technical",
        csf_function="protect",
        reference_control=policy_reference,
    )
    user = User.objects.create_user(
        email=f"bulk-no-promotion-{uuid.uuid4().hex}@tests.invalid"
    )
    _grant_permissions(
        user,
        world["folder"],
        "view_referencecontrol",
        "add_referencecontrol",
        "view_appliedcontrol",
        "change_appliedcontrol",
    )
    client = APIClient()
    client.force_authenticate(user)

    preview = client.get(
        f"/api/reference-controls/{policy_reference.id}/syncable-applied-controls/"
    )
    sync = client.post(
        f"/api/reference-controls/{policy_reference.id}/sync-applied-controls/",
        {},
        format="json",
    )

    assert preview.status_code == 200, preview.content
    assert sync.status_code == 200, sync.content
    assert preview.json() == []
    assert sync.json() == []
    ordinary_control.refresh_from_db()
    assert ordinary_control.category == "technical"
    assert ordinary_control.csf_function == "protect"


def test_applied_control_hidden_only_relationship_does_not_set_linked_flag(
    relationship_world,
):
    world = relationship_world
    target = _target(AppliedControl, world["folder"])
    target.requirement_assessments.add(world["unassigned"])
    client = APIClient()
    client.force_authenticate(world["respondent"])

    response = client.get(f"/api/applied-controls/{target.id}/")

    assert response.status_code == 200, response.content
    payload = response.json()
    assert payload["requirement_assessments"] == []
    assert "requirement_assessments" not in payload["linked_models"]


@pytest.mark.parametrize("_policy_field,target_model,serializer_class", TARGETS)
def test_governed_request_rejects_mixed_many_relationship_before_save(
    relationship_world,
    _policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    actor = Actor.objects.get(user=world["respondent"])
    owner_field = "owners" if target_model is SecurityException else "owner"
    serializer = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
        **{owner_field: [str(actor.id)]},
    )

    assert not serializer.is_valid()
    assert "requirement_assessments" in serializer.errors

    target.refresh_from_db()
    owner_manager = target.owners if target_model is SecurityException else target.owner
    assert not owner_manager.exists()
    assert set(target.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id
    }


def test_team_actor_does_not_confer_inverse_write_authority(relationship_world):
    world = relationship_world
    team = Team.objects.create(name="unlocked membership", folder=world["folder"])
    team.members.add(world["respondent"])
    team_actor = Actor.objects.get(team=team)
    world["assignment"].actor.clear()
    world["assignment"].actor.add(team_actor)
    target = _target(Evidence, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])

    serializer = _serializer(
        EvidenceWriteSerializer,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_a"],),
    )

    with pytest.raises(PermissionDenied):
        serializer.is_valid(raise_exception=True)

    client = APIClient()
    client.force_authenticate(world["respondent"])
    detail = client.get(f"/api/evidences/{target.id}/")
    filtered = client.get(
        "/api/evidences/",
        {"requirement_assessments": str(world["assigned_a"].id)},
    )

    assert detail.status_code == 200, detail.content
    assert _represented_ids(detail.json()["requirement_assessments"]) == {
        str(world["assigned_a"].id)
    }
    assert filtered.status_code == 200, filtered.content
    rows = filtered.json().get("results", filtered.json())
    assert str(target.id) in {row["id"] for row in rows}


@pytest.mark.parametrize("policy_field,target_model,serializer_class", CONCRETE_TARGETS)
@pytest.mark.parametrize("hidden_parent", ("requirement_node", "framework"))
def test_hidden_parent_is_excluded_before_relationship_stringification(
    relationship_world,
    monkeypatch,
    policy_field,
    target_model,
    serializer_class,
    hidden_parent,
):
    world = relationship_world
    row = world["assigned_a"]
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(row)
    original_viewable = RoleAssignment.get_viewable_object_ids

    def scoped_viewable(user, model):
        ids = original_viewable(user, model)
        if hidden_parent == "requirement_node" and model is RequirementNode:
            return [item for item in ids if item != row.requirement_id]
        if hidden_parent == "framework" and model is Framework:
            return [item for item in ids if item != world["assessment"].framework_id]
        return ids

    original_string = RequirementAssessment.__str__

    def guarded_string(instance):
        if instance.id == row.id:
            raise AssertionError("hidden RequirementAssessment was stringified")
        return original_string(instance)

    monkeypatch.setattr(RoleAssignment, "get_viewable_object_ids", scoped_viewable)
    monkeypatch.setattr(RequirementAssessment, "__str__", guarded_string)
    read = READ_SERIALIZERS[target_model](
        target, context={"request": _request_for(world["respondent"])}
    )

    assert read.data["requirement_assessments"] == []
    write = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(row,),
    )
    with pytest.raises(PermissionDenied):
        write.is_valid(raise_exception=True)
    assert policy_field in world["assessment"].field_visibility


@pytest.mark.parametrize("policy_field,target_model,serializer_class", CONCRETE_TARGETS)
def test_cross_framework_parent_chain_is_hidden_and_rejected(
    relationship_world,
    monkeypatch,
    policy_field,
    target_model,
    serializer_class,
):
    world = relationship_world
    other_framework = Framework.objects.create(
        name="Other framework",
        urn=f"urn:test:inverse-other:{uuid.uuid4().hex}",
        ref_id="OTHER",
        folder=world["folder"],
        min_score=0,
        max_score=4,
    )
    other_node = RequirementNode.objects.create(
        name="Cross framework",
        urn=f"{other_framework.urn}:node",
        ref_id="CROSS",
        framework=other_framework,
        folder=world["folder"],
        assessable=True,
    )
    row = world["assigned_a"]
    RequirementAssessment.objects.filter(id=row.id).update(requirement=other_node)
    row.refresh_from_db()
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(row)

    def forbidden_string(_instance):
        raise AssertionError("corrupt RequirementAssessment was stringified")

    monkeypatch.setattr(RequirementAssessment, "__str__", forbidden_string)
    read = READ_SERIALIZERS[target_model](
        target, context={"request": _request_for(world["respondent"])}
    )
    assert read.data["requirement_assessments"] == []

    write = _serializer(
        serializer_class,
        target=target,
        user=world["respondent"],
        requested=(row,),
    )
    with pytest.raises(PermissionDenied):
        write.is_valid(raise_exception=True)
    assert policy_field in world["assessment"].field_visibility


@pytest.mark.parametrize("policy_field,target_model,_serializer_class", TARGETS)
def test_requirement_assessment_filter_hides_missing_and_invisible_operands(
    relationship_world,
    policy_field,
    target_model,
    _serializer_class,
):
    world = relationship_world
    target = _target(target_model, world["folder"])
    target.requirement_assessments.add(world["assigned_a"], world["unassigned"])
    client = APIClient()
    client.force_authenticate(world["respondent"])
    endpoint = f"/api/{ENDPOINTS[target_model]}/"

    visible = client.get(endpoint, {"requirement_assessments": world["assigned_a"].id})
    hidden = client.get(endpoint, {"requirement_assessments": world["unassigned"].id})
    missing = client.get(endpoint, {"requirement_assessments": uuid.uuid4()})

    assert visible.status_code == 200, visible.content
    assert hidden.status_code == missing.status_code == 403
    assert hidden.json() == missing.json()
    assert policy_field in world["assessment"].field_visibility


def test_evidence_create_rolls_back_when_revision_creation_fails(
    relationship_world,
    monkeypatch,
):
    world = relationship_world
    name = f"rollback-evidence-{uuid.uuid4().hex}"

    def fail_revision(*_args, **_kwargs):
        raise RuntimeError("revision failure")

    monkeypatch.setattr(EvidenceRevision.objects, "get_or_create", fail_revision)
    serializer = EvidenceWriteSerializer(
        data={
            "name": name,
            "folder": str(world["folder"].id),
        },
        context={"request": _request_for(world["respondent"])},
    )
    serializer.is_valid(raise_exception=True)

    with pytest.raises(RuntimeError, match="revision failure"):
        serializer.save()
    assert not Evidence.objects.filter(name=name).exists()


def test_batch_conflict_preserves_detail_and_code(relationship_world, monkeypatch):
    world = relationship_world
    target = _target(Evidence, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])

    def fail_final_reproof(_context, _plan):
        raise RequirementAssessmentRelationshipConflict()

    monkeypatch.setattr(
        "core.requirement_assessment_relationships._revalidate_authority_snapshot",
        fail_final_reproof,
    )
    client = APIClient()
    client.force_authenticate(world["respondent"])
    response = client.post(
        "/api/evidences/batch-action/",
        {
            "action": "add_m2m",
            "ids": [str(target.id)],
            "field": "requirement_assessments",
            "value": [str(world["assigned_b"].id)],
        },
        format="json",
    )

    assert response.status_code == 200, response.content
    failure = response.json()["failed"][0]
    assert failure["code"] == "requirement_assessment_relationship_conflict"
    assert "changed" in str(failure["error"])


def test_governed_webhook_failure_does_not_reverse_committed_api_update(
    relationship_world,
    monkeypatch,
    django_capture_on_commit_callbacks,
):
    world = relationship_world
    target = _target(Evidence, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])

    def fail_webhook(*_args, **_kwargs):
        raise RuntimeError("webhook unavailable")

    monkeypatch.setattr("core.views.dispatch_webhook_event", fail_webhook)
    client = APIClient()
    client.force_authenticate(world["respondent"])
    with django_capture_on_commit_callbacks(execute=True):
        response = client.patch(
            f"/api/evidences/{target.id}/",
            {"requirement_assessments": [str(world["assigned_b"].id)]},
            format="json",
        )

    assert response.status_code == 200, response.content
    assert set(target.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_b"].id
    }


def test_governed_api_create_commits_target_and_relationship_before_webhook(
    relationship_world,
    monkeypatch,
    django_capture_on_commit_callbacks,
):
    world = relationship_world
    webhook_events = []

    def record_webhook(instance, event_type, *, serializer):
        webhook_events.append((instance.id, event_type, serializer.instance.id))
        assert set(instance.requirement_assessments.values_list("id", flat=True)) == {
            world["assigned_a"].id
        }

    monkeypatch.setattr("core.views.dispatch_webhook_event", record_webhook)
    client = APIClient()
    client.force_authenticate(world["respondent"])
    with django_capture_on_commit_callbacks(execute=True):
        response = client.post(
            "/api/evidences/",
            {
                "name": f"atomic-api-create-{uuid.uuid4().hex}",
                "folder": str(world["folder"].id),
                "requirement_assessments": [str(world["assigned_a"].id)],
            },
            format="json",
        )

    assert response.status_code == 201, response.content
    created = Evidence.objects.get(id=response.json()["id"])
    assert set(created.requirement_assessments.values_list("id", flat=True)) == {
        world["assigned_a"].id
    }
    assert webhook_events == [(created.id, "created", created.id)]


def test_applied_control_integration_failure_does_not_block_governed_webhook(
    relationship_world,
    monkeypatch,
    django_capture_on_commit_callbacks,
):
    world = relationship_world
    target = _target(AppliedControl, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    root = Folder.get_root_folder()
    provider = IntegrationProvider.objects.create(
        name=f"inverse-ra-itsm-{uuid.uuid4().hex}",
        provider_type=IntegrationProvider.ProviderType.ITSM,
        folder=root,
    )
    IntegrationConfiguration.objects.create(
        provider=provider,
        folder=root,
        credentials={},
        settings={
            "models": {
                "applied_control": {
                    "table_name": "x_governed_control",
                    "field_map": {"name": "short_description"},
                }
            }
        },
        webhook_secret="test-only-secret",
    )
    callback_order = []

    def fail_integration_schedule(*_args, **_kwargs):
        callback_order.append("integration")
        raise RuntimeError("scheduler unavailable")

    def record_governed_webhook(instance, event_type, *, serializer):
        callback_order.append("webhook")
        assert instance.id == target.id
        assert event_type == "updated"
        assert serializer.instance.id == target.id

    # An RA-only request deliberately cannot carry an integration-syncable
    # scalar.  Force only the model's change-detector result so this request
    # exercises the real save -> integration lookup -> on_commit callback path;
    # callback registration and execution remain production behavior.
    monkeypatch.setattr(
        AppliedControl,
        "_capture_sync_changed_fields",
        lambda _self: ["name"],
    )
    monkeypatch.setattr(
        "integrations.tasks.sync_object_to_integrations.schedule",
        fail_integration_schedule,
    )
    monkeypatch.setattr("core.views.dispatch_webhook_event", record_governed_webhook)
    client = APIClient()
    client.force_authenticate(world["respondent"])

    with django_capture_on_commit_callbacks(execute=True):
        response = client.patch(
            f"/api/applied-controls/{target.id}/",
            {"requirement_assessments": [str(world["assigned_b"].id)]},
            format="json",
        )

    assert response.status_code == 200, response.content
    assert set(
        AppliedControl.objects.get(id=target.id).requirement_assessments.values_list(
            "id", flat=True
        )
    ) == {world["assigned_b"].id}
    assert callback_order == ["integration", "webhook"]


def test_security_status_notification_is_not_scheduled_for_rejected_mixed_payload(
    relationship_world,
    monkeypatch,
):
    world = relationship_world
    target = _target(SecurityException, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    actor = Actor.objects.get(user=world["respondent"])
    target.owners.add(actor)
    queued = []

    monkeypatch.setattr(
        "core.tasks.send_security_exception_status_notification",
        lambda *args: queued.append(args),
    )

    serializer = _serializer(
        SecurityExceptionWriteSerializer,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
        status=SecurityException.Status.IN_REVIEW,
    )
    assert not serializer.is_valid()
    assert list(serializer.errors) == ["requirement_assessments"]
    target.refresh_from_db()
    assert target.status == SecurityException.Status.DRAFT
    assert queued == []


def test_inverse_relationship_uses_bounded_mail_compatible_lock_order(
    relationship_world,
    monkeypatch,
):
    """Pin this boundary's order and the shared mail subsequence only."""

    world = relationship_world
    target = _target(Evidence, world["folder"])
    target.requirement_assessments.add(world["assigned_a"])
    serializer = _serializer(
        EvidenceWriteSerializer,
        target=target,
        user=world["respondent"],
        requested=(world["assigned_b"],),
    )
    serializer.is_valid(raise_exception=True)

    locked_models: list[type] = []
    original_select_for_update = QuerySet.select_for_update

    def tracked_select_for_update(queryset, *args, **kwargs):
        locked_models.append(queryset.model)
        return original_select_for_update(queryset, *args, **kwargs)

    monkeypatch.setattr(QuerySet, "select_for_update", tracked_select_for_update)
    serializer.save()

    expected_order = [
        Framework,
        ComplianceAssessment,
        RequirementAssignment,
        RequirementNode,
        RequirementAssessment,
        Evidence,
        Evidence.requirement_assessments.through,
        RequirementAssignment.requirement_assessments.through,
        RequirementAssignment.actor.through,
        Actor,
        User,
    ]
    assert locked_models == expected_order
    shared_suffix = [RequirementAssignment.actor.through, Actor, User]
    assert [model for model in locked_models if model in shared_suffix] == shared_suffix
